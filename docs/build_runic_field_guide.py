from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path(r"E:\Valheim Mods\ChazmanModsRepo\docs\Runic_Mod_Testing_Primer_and_Reference.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(36, 42, 48)
MUTED = RGBColor(92, 102, 112)
PALE_BLUE = "E8EEF5"
PALE_GOLD = "FFF4D6"
PALE_GREEN = "EAF4E8"
PALE_RED = "F9E9E7"
WHITE = RGBColor(255, 255, 255)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=110, bottom=70, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_fixed_table(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    total = sum(widths)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(total))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "110")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[idx]))
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run(run, size=None, bold=None, color=None, italic=None, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def add_field(paragraph, field):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def configure_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(9.1)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.05

    for name, size, before, after, color in (
        ("Title", 25, 0, 5, DARK_BLUE),
        ("Subtitle", 11.5, 0, 8, MUTED),
        ("Heading 1", 15.5, 10, 6, BLUE),
        ("Heading 2", 11.5, 7, 3, BLUE),
        ("Heading 3", 9.8, 5, 2, DARK_BLUE),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(8.9)
        style.paragraph_format.left_indent = Inches(0.32)
        style.paragraph_format.first_line_indent = Inches(-0.16)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.03

    if "Checklist" not in styles:
        cstyle = styles.add_style("Checklist", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cstyle = styles["Checklist"]
    cstyle.font.name = "Calibri"
    cstyle.font.size = Pt(8.9)
    cstyle.font.color.rgb = INK
    cstyle.paragraph_format.left_indent = Inches(0.22)
    cstyle.paragraph_format.first_line_indent = Inches(-0.22)
    cstyle.paragraph_format.space_after = Pt(2.2)
    cstyle.paragraph_format.line_spacing = 1.03

    if "Compact" not in styles:
        compact = styles.add_style("Compact", WD_STYLE_TYPE.PARAGRAPH)
    else:
        compact = styles["Compact"]
    compact.font.name = "Calibri"
    compact.font.size = Pt(8.4)
    compact.font.color.rgb = INK
    compact.paragraph_format.space_after = Pt(2)
    compact.paragraph_format.line_spacing = 1.0

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("RUNIC MOD SUITE  |  IN-GAME TEST PRIMER & REFERENCE")
    set_run(r, 7.5, True, MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    r = fp.add_run("Valheim 0.221.12  •  Page ")
    set_run(r, 7.5, False, MUTED)
    add_field(fp, "PAGE")


def title_block(doc, kicker, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(kicker.upper())
    set_run(r, 8.1, True, BLUE)
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_run(r, 23, True, DARK_BLUE)
    p = doc.add_paragraph(style="Subtitle")
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(subtitle)
    set_run(r, 10.6, False, MUTED, True)


def heading(doc, text, level=2):
    return doc.add_paragraph(text, style=f"Heading {level}")


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.47)
    p.add_run(text)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def check(doc, text):
    p = doc.add_paragraph(style="Checklist")
    r = p.add_run("☐  ")
    set_run(r, 10, True, BLUE, font="DejaVu Sans")
    p.add_run(text)
    return p


def callout(doc, label, text, fill=PALE_GOLD):
    table = doc.add_table(rows=1, cols=1)
    set_fixed_table(table, [10166])
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + ": ")
    set_run(r, 8.8, True, DARK_BLUE)
    r = p.add_run(text)
    set_run(r, 8.8, False, INK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.line_spacing = 0.15


def small_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_fill(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run(r, 8.2, True, DARK_BLUE)
    set_repeat_table_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_run(r, 8.0, False, INK)
    set_fixed_table(table, widths)
    return table


def page_break(doc):
    doc.add_page_break()


def add_primer_page_1(doc):
    title_block(doc, "Three-page field test • Page 1 of 3", "Runic Suite In-Game Primer", "A single practical pass for the 20 packaged mods. Mark PASS only when the observed result matches the expected result.")
    callout(doc, "Before you start", "Back up the character and world. Use one known profile with all 20 packages enabled. Remove Build Camera Custom Hammers Edition before testing Runic Build Camera. Keep LogOutput.log from this session if anything fails.")

    heading(doc, "1  Startup and foundation health", 2)
    check(doc, "Launch Valheim, reach the main menu, enter a mapped test world, and wait for the local player to spawn. Expected: no Runic fatal/error popup and ordinary movement/UI remain available.")
    check(doc, "On a dedicated server, confirm the same required Runic versions are installed on server and client. If account-bound features are used, verify the administrator enrolled the connected character with runic_bind.")
    check(doc, "Open BepInEx LogOutput.log after the session only if needed. Core, Persistence, Permissions, Transactions, Sentinel, Velocity, and World Engine should report ready/monitoring rather than a compatibility stop.")

    heading(doc, "2  Inventory", 2)
    check(doc, "Open inventory. Expected: the native bottom row is subtly highlighted and labeled Helmet, Chest, Legs, Cape, Utility, Quick 1, Quick 2, Quick 3; no separate gameplay quick-slot HUD appears.")
    check(doc, "With one equipment role empty, pick up or craft that recognized gear type. Expected: it moves to the exact role and equips. Craft/pick up a second piece for the occupied role. Expected: it stays in ordinary backpack space.")
    check(doc, "Manually equip the replacement. Expected: it enters the role and the old item swaps into the replacement's original slot; no item disappears or duplicates.")
    check(doc, "Put an allowed consumable/tool/utility item in Quick 1 and press Left Alt+1. Repeat for Quick 2/3 with Alt+2/3. Expected: each activates manually and vanilla hotbar input does not also fire.")
    check(doc, "Focus an ordinary slot and press Left Alt+L to lock it. Press Left Alt+I with inventory open. Expected: safe general rows sort, while hotbar, equipment/quick row, equipped items, and locked coordinates stay protected.")

    heading(doc, "3  Storage and Interaction", 2)
    check(doc, "Hover a closed, authorized, non-personal chest within reach. Expected: a bounded Contents summary appears. Open/in-use, private, denied, distant, or unsynchronized chests show vanilla hover only.")
    check(doc, "Place matching items in nearby public chests, keep matching stacks on the player, then press Left Alt+Q with inventory closed or player-only inventory open. Expected: Quick Stack moves only eligible matching items.")
    check(doc, "Open one authorized chest and press Left Alt+A. Expected: Store All fills that exact chest with eligible carried items but leaves equipped, quest, hotbar-protected, locked, and special-row items carried.")
    check(doc, "Test Left Alt+R Restock, Alt+F Search, opened-chest Alt+S Sort, and Alt+C Consolidate. Expected: each reports a clear result/no-op reason and preserves protected items.")
    check(doc, "Hold Use on a supported smelter/cooking/fermenter/shield input. Expected: vanilla-rate repeat. In an open chest, Alt+left-click a stack. Expected: guarded full-stack transfer.")


def add_primer_page_2(doc):
    title_block(doc, "Three-page field test • Page 2 of 3", "Build, Farm, Craft & Automate", "Use cheap materials and disposable test pieces. Vanilla costs, wards, range, station, and placement rules should remain authoritative.")
    heading(doc, "4  Build Camera and Precision Build Tool", 2)
    check(doc, "Equip hammer, select a piece, and press B. Move/look with normal controls; jump/crouch move vertically and Run speeds up. Expected: camera remains bounded, avatar stays in place and vulnerable, and normal placement/removal rules still apply.")
    check(doc, "Exit camera, select a piece, press P for Precision mode. Test wheel yaw, Alt+wheel pitch, Shift+wheel roll, V fine control, Alt+arrows/Page keys movement, and F10 reset. Expected: readout changes and normal placement still validates/costs resources.")
    check(doc, "Aim at a placed piece and try Keypad0 rotation match, Keypad7 position, Keypad8 full transform, then KeypadPeriod after two placements. Expected: only transforms copy; prefab/ownership/health do not.")
    check(doc, "While the piece menu is open, try F6 search cycle, F7 favorite, F8 favorite cycle, F9 recent. In Precision mode with menu closed, cautiously test F11 undo and F12 bounded area repair on disposable pieces.")

    heading(doc, "5  Agriculture", 2)
    check(doc, "Equip cultivator and select a crop. Expected: pattern ghosts and the bottom Agriculture bar replace only vanilla BuildHints. Use Alt+right-click or Alt+O to cycle; Numpad1–7 selects Row/Grid/Circle/Star/Triangle/Half Circle/Trapezoid.")
    check(doc, "Use Alt+wheel rows, Shift+wheel columns, Alt+Shift+wheel spacing, then Alt+P. Expected: only valid affordable positions plant, each normal cost/stamina/durability charge applies, and invalid positions remain skipped.")
    check(doc, "Aim at an eligible Pickable and press Alt+E. Expected: exact same-prefab nearby items harvest within the configured radius/cap. With a matching grown crop, select that crop and confirm replant with Alt+T.")

    heading(doc, "6  Crafting", 2)
    check(doc, "Put recipe/build materials in eligible nearby public chests. In solo, listen-server, or dedicated-client play, open a station. Expected: requirement rows show carried/nearby/total/missing and normal Craft/Build consumes player inventory first, then eligible nearby materials.")
    check(doc, "Press Repair once with several repairable items. Expected: Repair All repairs every item Valheim currently considers repairable. On a dedicated server, also verify one remote craft, upgrade, and build from an eligible server chest.")

    heading(doc, "7  Production", 2)
    check(doc, "At a smelter/kiln, Alt+Use the input/fuel/output control, then Alt+Use a target chest within 30 seconds. Add ore/fuel. Expected: linked roles move complete vanilla inputs/outputs without changing timers, ratios, capacity, roof/fire, or fuel rules.")
    check(doc, "For a cooking/recipe/fermenter station, link Input and either throughput Output or the Replenishment manager. Put a physical example of each desired product in a replenishment chest and refresh it. Expected: complete batches route only to authorized linked destinations.")
    check(doc, "Remove the last example or fill the destination. Expected: new replenishment pauses; existing in-flight assignments remain bound. Test configured ingredient reserves if used.")

    heading(doc, "8  Display Stands", 2)
    check(doc, "Interact with a configured item stand/armor stand. Expected: compact native container UI; normal drag preserves full item metadata. Hold Left Alt while interacting for one-item take/armor-container access, and use an armor stand switch to swap the supported loadout.")


def add_primer_page_3(doc):
    title_block(doc, "Three-page field test • Page 3 of 3", "World, Map, Safety & Finish", "Complete the map/portal and display checks, then preserve only the evidence needed for any exact failure.")
    heading(doc, "9  Portals and Groups", 2)
    check(doc, "Place two vanilla portals with the same tag. Expected: Standard Pair behaves exactly like vanilla and pairs only with another Standard Pair.")
    check(doc, "Aim at an unconnected portal and press Alternate Place+Use (default Left Shift+E). Enter network|public|TEST|Home|both; configure another endpoint with the same case-sensitive TEST network and a unique name. Expected: network animation activates.")
    check(doc, "Walk into a depart/both network portal. Expected: player is protected/held until owner acknowledgement, then the map opens with authorized same-network arrivals. Click destination or Esc.")
    check(doc, "Open the normal large map anywhere and press unmodified P. Expected: world-wide portal directory toggles and shows authorized [Vanilla], [Public], your [Private], and your [Group] portals without requiring proximity.")
    check(doc, "Optional group check: use chat /group create Test Crew, invite/accept another connected player, /group use Test Crew, then create a group portal. Expected: only current members see/use it.")

    heading(doc, "10  Exploration and Awareness", 2)
    check(doc, "Create saved pins on explored map cells, including one named [boat] Test. Open the large map. Expected: Exploration panel appears. Type in Search, change Category/Source, click a result; map centers and shows LAST KNOWN distance/direction. Escape should release search focus, not permanently disable the panel.")
    check(doc, "Control a ship and reopen the map. Expected: LIVE LOCAL SHIP readout. Controller keeps vanilla map controls; browser selection uses mouse/keyboard.")
    check(doc, "Return to gameplay. Expected: Awareness shows enabled local food/effect/comfort/context panels and yields completely while inventory, map, menus, chat, or another interactive UI is open.")

    heading(doc, "11  Safety", 2)
    check(doc, "Try hammer-removing an occupied chest or ship/cart. Expected: first identical action warns; the second on the unchanged target within the configured window proceeds through vanilla checks.")
    check(doc, "Try overwriting a non-empty portal tag or sacrificing a configured rare item. Expected: contextual confirmation; locked/equipped/quest items remain denied. Do not use valuable items for this test.")

    heading(doc, "12  Silent foundations and diagnostics", 2)
    check(doc, "Velocity: startup completes and its local manifest cache/timeline reports ready in the log; no player UI is expected.")
    check(doc, "Sentinel: first-run Optional mode records/monitors policy state without locking the group out; no gameplay UI is expected.")
    check(doc, "World Engine: aggregate observation/ownership registry reports ready; no world mutation or UI is expected.")
    check(doc, "Core, Persistence, Permissions, and Transactions are indirectly verified when Storage dedicated actions, Portals/groups, Safety admission, and Production links complete without provider/protocol errors.")

    heading(doc, "Session verdict", 2)
    small_table(doc, ["Result", "Record"], [
        ("Overall", "☐ PASS   ☐ FAIL"),
        ("Profile / world", "____________________________________________"),
        ("Date / Valheim", "____________________________________________"),
        ("First exact failure", "____________________________________________"),
        ("Evidence saved", "☐ screenshot   ☐ LogOutput.log   ☐ reproduction steps"),
    ], [2250, 7916])


MODULES = [
    {
        "name": "Runic Core", "version": "1.0.0",
        "purpose": "Non-gameplay suite foundation for module discovery, typed services, input-conflict reporting, and actionable notifications.",
        "features": [
            "Registers immutable module identity, semantic/protocol version, capabilities, and services.",
            "Publishes canonical capability IDs and provider-neutral inventory durability contracts.",
            "Detects exact keybinding chord conflicts without silently remapping controls.",
            "Rate-limits actionable player notifications by module, reason, and context.",
            "Declares whether a setting is local, server-authoritative, or world-authoritative; it does not synchronize another mod's config.",
        ],
        "use": ["No player controls or world behavior.", "Keep it installed before every Runic package that declares it.", "Optional diagnostics: enable VerboseLogging only when investigating registry transitions."],
        "verify": ["Main menu and world load without a Core protocol error.", "Dependent mods publish ready messages/services.", "A reported Runic denial contains both a reason and a remedy."],
        "limits": ["No Harmony patches, inventory mutation, world data, RPC gameplay, or save changes.", "A missing/incompatible service is intentionally surfaced to the dependent module."],
        "deps": "BepInExPack Valheim 5.4.2333"
    },
    {
        "name": "Runic Persistence", "version": "1.0.0",
        "purpose": "Shared atomic migration, direct connection-bound RPC, compatibility admission, and account-to-character binding foundation.",
        "features": [
            "Owns canonical runic.<module>.<key> names and forward-only migration chains.",
            "Preserves unknown data and atomically replaces validated snapshots.",
            "Provides bounded direct ZRpc requests without trusting client-supplied routed sender IDs.",
            "Negotiates per-capability protocols and compatibility claims before world admission.",
            "Maintains world-scoped, reverse-unique backend-account to Valheim-player bindings for features that explicitly require AccountBoundPlayer.",
        ],
        "use": ["Ordinary play has no UI.", "Dedicated Steam admin: open F5, run runic_bind peers, then runic_bind enroll <peerUid> for each character that needs group/private/ward owner features.", "Use runic_bind list [offset] to verify enrollment."],
        "verify": ["Matching client joins without protocol/admission failure.", "runic_bind list shows the expected world-specific pairs.", "Portals/groups/remote Storage actions that require actor binding succeed for the enrolled player."],
        "limits": ["Enrollment is explicit and world-specific; display names and payload IDs are not authority.", "No recipes, pieces, items, or standalone gameplay."],
        "deps": "BepInExPack; Runic Core 1.0.0"
    },
    {
        "name": "Runic Permissions", "version": "1.0.0",
        "purpose": "Canonical identities, ownership, action scopes, ward-aware policy evaluation, and server-owned player Groups.",
        "features": [
            "Defines Everyone, Approved, Owner, Nobody, Ward, Ward+exceptions, and Group policies.",
            "Evaluates each action independently; container discovery never implies withdraw or configuration permission.",
            "Publishes permission, group-membership, and active-group services.",
            "Stores world-scoped Group catalog and current active Group selection.",
            "Optional admin bypass is off by default and cannot allow unless audit succeeds.",
        ],
        "use": ["In normal chat: /group create <name>; /group invite <connected name>; invitee uses /group accept <name>.", "Use /group list, /group use <name>, /group active, /group members, /group role, /group remove, /group transfer, /group rename, /group leave, or /group delete.", "Group commands are private local chat commands, not public messages."],
        "verify": ["Create a Group and invite one connected player.", "Invitee accepts and both /group members results agree.", "A Group portal is visible to a member and hidden from a nonmember."],
        "limits": ["No ward discovery or direct gameplay mutation; the consuming mod supplies current ward/object evidence.", "Ambiguous identity, membership, or overlapping hostile ward evidence denies."],
        "deps": "BepInExPack; Runic Core; Runic Persistence 1.0.0"
    },
    {
        "name": "Runic Transactions", "version": "1.0.0",
        "purpose": "All-or-nothing resource transactions, persistent world-object identity, and restart-durable claims used by trusted Runic adapters.",
        "features": [
            "Canonicalizes exact resource allocations and acquires endpoint locks in stable order.",
            "Revalidates permission and balances before commit; retry is idempotent.",
            "Provides process-wide non-reentrant mutation gate for short synchronous changes.",
            "Persists durable composite/world-object operations and exact endpoint claims.",
            "Blocks open, stack/take-all, auto-destroy, or destruction while unresolved evidence is required for recovery.",
            "Carries neutral container query/transfer contract types but does not advertise a Storage service itself.",
        ],
        "use": ["No direct player control or panel.", "Do not manually remove operation locks or journal evidence.", "Let the owning gameplay mod report and recover a pending action."],
        "verify": ["A Storage dedicated mutation or Production linked operation completes without duplicate consumption.", "If a durable operation is intentionally interrupted, the affected endpoint stays protected until recovery finishes.", "No transaction-provider startup error appears."],
        "limits": ["Infrastructure only; it does not scan the world or automate containers by itself.", "Uncertain durable evidence fails closed instead of guessing."],
        "deps": "BepInExPack; Runic Core; Runic Persistence 1.0.0"
    },
    {
        "name": "Runic Build Camera", "version": "1.0.0",
        "purpose": "Bounded detached camera for ordinary Valheim building, with local loose-item pickup and Wisplight support.",
        "features": [
            "Detached camera movement/look while the avatar remains in normal building state.",
            "Default 60 m camera range and 100 m avatar-to-target remote-action cap; camera ray remains 50 m.",
            "Placement, repair, and removal continue through vanilla costs, stations, wards, collisions, and server behavior.",
            "Optional bounded loose-world-item pickup near camera; never opens or searches containers.",
            "Optional movement of the player's active Wisplight demister to camera with range multiplier.",
            "Configurable speed, fast multiplier, movement frame, and mouse/controller inversion.",
        ],
        "use": ["Equip hammer and select a piece; press B to enter/leave.", "Use normal movement/look; jump/crouch move vertically; Run applies fast speed.", "Edit ToggleShortcut and bounded camera/action ranges in chazman.RunicBuildCamera.cfg."],
        "verify": ["B enters and exits without leaving the camera/effect detached.", "Place and remove a cheap test piece from camera; ordinary cost/ward/range rules still apply.", "With Wisplight and pickup enabled, verify the effect follows and only loose nearby drops are collected."],
        "limits": ["Avatar stays at origin and remains vulnerable.", "Do not install beside Build Camera Custom Hammers Edition; Precision Build Tool is compatible.", "Client-side caps are not server anti-cheat authority."],
        "deps": "BepInExPack only"
    },
    {
        "name": "Runic Display Stands", "version": "1.3.1",
        "purpose": "Turns configured item and armor stands into compact native container-style storage with loadout swapping.",
        "features": [
            "Item stands accept any item; armor stands accept one supported item per supplied attachment category.",
            "Native container UI supports normal drag, Take all, and Use/equip buttons.",
            "Armor stand switch swaps its loadout with equipped armor and first matching hotbar weapon/shield.",
            "Preserves quality, durability, crafter, variant, enchantment, and other custom item data.",
            "Invalid armor-slot items return to player or safely drop when inventory is full.",
            "Forsaken stones, boss holders, quest props, and other scripted stands keep original behavior.",
        ],
        "use": ["Interact with a configured stand to open it.", "Hold configured Take One key (Left Alt) while interacting to remove one item or access the armor-container route indicated by prompts.", "Interact with armor stand switch to swap loadout."],
        "verify": ["Store and retrieve a metadata-bearing item; its data remains unchanged.", "Load one item in each supported armor category; duplicate/utility mismatch is rejected safely.", "Swap the armor loadout twice and verify exact conservation."],
        "limits": ["All interacting players should install it.", "Supported stand prefab list is host-synchronized; only declared BepIn dependency."],
        "deps": "BepInExPack Valheim"
    },
    {
        "name": "Runic Inventory", "version": "1.0.0",
        "purpose": "Uses Valheim's native bottom row for five equipment roles and three manual quick slots without adding capacity or a second store.",
        "features": [
            "Labels native bottom cells Helmet, Chest, Legs, Cape, Utility, Quick 1–3 only while inventory is open.",
            "Auto-places/equips recognized gear when its exact role is empty; otherwise leaves new gear in ordinary space.",
            "Manual replacement performs a lossless exact-position swap with the prior role occupant.",
            "Quick 1–3 accept supported consumables, tools, and utility items and are manual-use only.",
            "Coordinate slot locks protect explicit items; every occupied special-row role is protected from Storage transfers.",
            "Safe sort permutes configured general rows only; never merges stacks.",
            "Pickup preview reports safe capacity/projected weight; optional exact pickup deny list is non-destructive.",
        ],
        "use": ["Alt+1/2/3 use Quick 1/2/3.", "With inventory open, Alt+I sorts configured safe rows.", "Focus/select a player slot, then Alt+L toggles its lock.", "Controller: hold JoyAltKeys; JoyMap/Y/RB use Quick 1–3, A sorts, B locks."],
        "verify": ["Bottom row labels appear; no external quick-slot HUD.", "Empty-role pickup equips; occupied-role pickup stays in backpack; manual replacement swaps.", "Alt quick use does not also activate vanilla hotbar.", "Locked/special-row items survive Storage Quick Stack/Store All and Inventory sort."],
        "limits": ["No added slots, weight, or stack size.", "Unknown/modded item categories stay ordinary rather than being guessed.", "If role evidence is incompatible, correct the reported cell/equipment; other moves fail closed."],
        "deps": "BepInExPack; Runic Core; Persistence; Transactions 1.0.0"
    },
    {
        "name": "Runic Storage", "version": "1.0.0",
        "purpose": "Authorized chest hover, Quick Stack, Store All, Restock, Search, Sort, and player-stack consolidation.",
        "features": [
            "Closed authorized chest hover lists bounded contents without opening or taking ownership.",
            "Quick Stack deposits eligible carried items only where matching items already exist.",
            "Store All deposits every eligible fitting carried item into the exact currently open/owned chest, including dedicated clients.",
            "Restock pulls configured targets; Search reports total and nearest direction; opened chest Sort orders items.",
            "Consolidate merges only serialization-compatible carried stacks.",
            "Respects equipped, quest, protected hotbar, Runic Inventory role/lock, access, ward, busy, claim, and quarantine gates.",
            "Publishes bounded container.query for Crafting and other optional peers.",
        ],
        "use": ["Alt+Q Quick Stack; Alt+R Restock; Alt+F Search; open chest + Alt+S Sort; Alt+A Store All; Alt+C Consolidate.", "Controller: hold JoyAltKeys; D-pad Down/Up/Right/Left = Quick/Restock/Search/Consolidate; R-stick = Sort.", "Set Restock targets and Search item in config."],
        "verify": ["Nonempty authorized closed chest shows contents; busy/private/distant chest does not.", "Quick Stack works with player inventory open or closed and skips protected items.", "Store All works in solo/listen and dedicated client on currently open owned chest.", "Restock/Search/Sort/Consolidate report clear outcome/no-op reasons."],
        "limits": ["Personal containers are excluded from nearby discovery.", "Dedicated Quick Stack/Restock/Sort require matching server modules and actor binding; uncertainty makes no change.", "Store All is exact-opened-container local ownership, not the remote multi-container saga."],
        "deps": "BepInExPack; Core; Persistence; Transactions; Inventory 1.0.0"
    },
    {
        "name": "Runic Interaction", "version": "1.0.0",
        "purpose": "Reversible interaction conveniences: hold-repeat, guarded transfers, equipment restore, menu memory, text validation, pickup filters, and optional doors.",
        "features": [
            "Hold Use repeats supported station input/fuel at vanilla 0.2 s cadence.",
            "Alt+left-click adds a guarded full-stack transfer; vanilla Ctrl+click and LT+X remain supported.",
            "Restores legal pre-swim hands or combat hands displaced by a temporary Tool when safe.",
            "Remembers crafting recipe/group per station prefab and craft/upgrade mode for the session.",
            "Validates portal/sign/tame text immediately before vanilla commit.",
            "Optional exact pickup filter leaves denied drops in world; Alt+Use bypasses intentionally.",
            "Authenticated delayed door auto-close exists but is disabled by default.",
        ],
        "use": ["Hold Use on supported smelter/cooking/empty-fermenter/shield inputs.", "Open a chest and Alt+left-click for full stack.", "Configure Pickup Filter.Items; hold Alt+Use or JoyRStick+Use to bypass.", "Enable AutoCloseDoors only if desired; default delay is 4 seconds."],
        "verify": ["Held station input repeats one normal action per cadence.", "Full-stack gesture obeys quest/lock/access/ward/drag rules.", "Swim or put away a temporary Tool; prior legal hands restore when safe.", "Filtered drop remains on ground; bypass picks it up."],
        "limits": ["No drag-sweep transfer or combat automation.", "Generic filter memory is unavailable; ordinary vanilla dragging remains.", "Door auto-close requires matching server/clients and enrollment."],
        "deps": "BepInExPack; Runic Core; Runic Persistence 1.0.0"
    },
    {
        "name": "Runic Crafting", "version": "1.0.0",
        "purpose": "Nearby-material crafting/building for solo, listen servers, and dedicated clients, plus Workshop Access and Repair All while preserving vanilla recipes and costs.",
        "features": [
            "Uses carried materials first, then eligible nearby containers for ordinary recipes.",
            "Supports station-gated building and an exact allowlist of audited stationless pieces.",
            "Shows carried/nearby/total/missing requirement counts and reason when nearby use is off.",
            "Preserves multi-craft amount, recipe knowledge, station level, output, quality, skill, placement, and complete costs.",
            "One Repair press repairs every item Valheim currently considers repairable.",
            "Independent Station Use and Local Material Use policies; nearby permission never grants chest opening.",
            "Uses Storage's optional bounded query when present but does not depend on Storage.",
        ],
        "use": ["No hotkey: use normal Craft/Build/Repair controls.", "Place eligible materials in nearby authorized chests and open the station/build menu; dedicated clients use the same controls.", "Station owner may use F5 runiccrafting_access show/station/materials/approve/unapprove commands."],
        "verify": ["Requirement rows accurately combine carried and nearby counts.", "Craft, upgrade, and build succeed with exact total cost and no partial loss on denial, including for a remote dedicated client.", "Repair button repairs all eligible items.", "A denied local-material player can still craft from carried materials."],
        "limits": ["Choose-one recipes, unknown stationless mod pieces, feast family, and recipe pinning remain vanilla/unavailable.", "Dedicated nearby use requires matching Core, Persistence, Permissions, Transactions, Inventory, and Crafting 1.0.0 on server and client."],
        "deps": "BepInExPack; Core; Persistence; Permissions; Transactions; Inventory 1.0.0"
    },
    {
        "name": "Runic Agriculture", "version": "1.0.0",
        "purpose": "Validated planting patterns, exact-Pickable area harvest, confirmed crop replant, and contextual crop/beehive information.",
        "features": [
            "Row, grid, filled circle, five-point star, right triangle, half circle, and tapered trapezoid previews.",
            "Per-position terrain, biome, cultivation, spacing, range, ward, no-build, dungeon, and player checks.",
            "Charges normal resources, stamina, and cultivator durability per successful plant; preview never auto-plants.",
            "Area harvest targets only the exact registered Pickable prefab within radius and 25-object hard cap.",
            "Replant offered only for an exact grown prefab of a registered Plant.",
            "Bottom Agriculture control bar replaces only BuildHints during crop preview and follows current input source.",
        ],
        "use": ["Alt+P confirm; Alt+right-click or Alt+O cycle; Numpad1–7 direct shape.", "Alt+wheel rows; Shift+wheel columns; Alt+Shift+wheel spacing; Alt+Shift+L side/mirror.", "Alt+E exact-Pickable harvest; Alt+T confirm matching replant.", "Controller crop preview: D-pad edits; Controller Alt+L-stick cycles; Alt+Place confirms; Alt+Use harvests."],
        "verify": ["Cultivator crop selection shows bounded ghosts and valid/total bar.", "Changing size/spacing updates preview and saved config; plain right-click still opens vanilla selector.", "Confirm plants only valid affordable positions with normal costs.", "Harvest ignores neighboring different prefabs and only offers valid crop replant."],
        "limits": ["No growth acceleration, extra yield, free seeds, relaxed biome rules, or unattended farming.", "Maximum 50 preview ghosts; harvest maximum 25.", "PickableItem treasure fixtures are excluded."],
        "deps": "BepInExPack; Core; Transactions; Persistence 1.0.0"
    },
    {
        "name": "Runic Production", "version": "1.0.0",
        "purpose": "Explicit chest-linked production lines for smelters, cooking, recipe stations, and fermenters with exemplar-driven replenishment.",
        "features": [
            "Smelter/kiln Input, optional Fuel, and Output links.",
            "Cooking/oven Input, optional Fuel, and throughput Output or multiple Replenishment chests.",
            "Cauldron/mead/prep Input plus multiple exemplar-driven Replenishment destinations.",
            "Fermenter Input and throughput Output or Replenishment destinations.",
            "Persistent station/chest identity, exact in-flight route assignments, round-robin destinations, complete-batch capacity checks.",
            "Per-prefab ingredient reserves and optional bounded nearby recipe staging (off by default).",
            "Preserves vanilla recipes, ratios, timers, fuel, roof/fire, capacity, and conversion rules.",
        ],
        "use": ["Alt+Use a station input/fuel/output control, then Alt+Use the target chest within 30 seconds.", "Cooking/fermenter: tap Ctrl+Alt+Use for throughput Output; hold 0.6 s for Replenishment manager.", "In manager, Alt+Use adds/refreshes a chest; Shift+Alt+Use removes that relation.", "Put at least one physical output example in each Replenishment chest and refresh it."],
        "verify": ["Linked smelter moves exact inputs/fuel/output and respects full/busy/ward/range stops.", "Replenishment makes only products physically represented and authorized in the chest.", "Removing final example pauses new work; re-adding/refreshing resumes.", "Restart/unload does not reroute in-flight cooking/fermenting output."],
        "limits": ["No recursive demand solver or generic unbounded base scan.", "Nearby ingredient staging is disabled by default.", "Recipe automation does not grant Cooking skill/bonus RNG/profile progress."],
        "deps": "BepInExPack; Core; Permissions; Transactions; Persistence 1.0.0"
    },
    {
        "name": "Runic Precision Build Tool", "version": "2.0.1",
        "purpose": "Explicit precision-mode six-degree-of-freedom preview control, transform matching, catalog helpers, conservative undo, and bounded area repair.",
        "features": [
            "Yaw/pitch/roll plus sway/heave/surge while Valheim owns final placement.",
            "Fine 1°/0.05 m controls, world/local movement frame, guides, per-axis/full reset.",
            "Copies aimed piece rotation, axes, world position, full transform, or validated snap side.",
            "Repeats relative transform from the last two successful placements.",
            "Search/favorites/recents over already-unlocked pieces only.",
            "Strict last-placement undo and bounded area repair with dedicated-server authorization.",
        ],
        "use": ["P toggles Precision mode. Wheel yaw; Alt+wheel pitch; Shift+wheel roll; V makes fine.", "Alt+arrows/Page Up/Down moves; hold G guides; F10 full reset.", "Keypad0 rotation; 1/2/3 axes; 4/5/6 XYZ; 7 position; 8 full; 9 snap; Decimal repeat; Shift+Keypad axis resets.", "Piece menu: F6 search, F7 favorite, F8 favorite cycle, F9 recent. Precision/menu closed: F11 undo, F12 repair."],
        "verify": ["Outside Precision mode, vanilla rotation/placement remains ordinary.", "Preview readout matches exact move/rotation and F10 resets.", "A match copies transform only, never prefab/owner/health.", "Undo/repair either completes exact safe target(s) or refuses without partial change."],
        "limits": ["No free resources or rule bypass.", "Undo is last eligible session placement only; repair is bounded and conservative.", "P overlaps the portal directory only in different contexts: build preview vs large map."],
        "deps": "BepInExPack; Core; Persistence; Transactions 1.0.0"
    },
    {
        "name": "Runic Portals", "version": "1.0.0",
        "purpose": "Vanilla Standard Pair preservation plus permission-aware public/private/group portal networks, walk-in routing, and a world-wide map directory.",
        "features": [
            "Standard portals pair and travel only with Standard portals exactly like vanilla.",
            "Network endpoints share case-sensitive NetworkName and unique display names with arrive/depart/both direction.",
            "Public, persisted-owner private, and current-membership Group visibility/authorization.",
            "Walk-in map picker shows authorized online same-network arrivals and protects the player/source portal while open.",
            "Normal large-map P directory shows authorized vanilla, public, private, and group portals anywhere in world without proximity.",
            "Network portal availability animation includes arrive-only endpoints.",
            "Bounded Return option and explicit one-way acknowledgement.",
        ],
        "use": ["Aim portal; normal Use edits vanilla tag; Alternate Place+Use (default Shift+E) opens Runic editor.", "Commands: network|public|NETWORK|NAME|both; private; or group after /group use <name>. Use arrive/depart as needed; standard restores vanilla mode.", "Walk into depart/both endpoint; click destination on picker or Esc.", "Open normal large map and press unmodified P to toggle world-wide directory."],
        "verify": ["Vanilla pair still works and never targets a Network endpoint.", "Two same-network Runic endpoints animate and route according to direction/policy.", "P directory works away from portals and labels [Vanilla]/[Public]/[Private]/[Group].", "Unauthorized/private/group-nonmember endpoints do not appear or travel."],
        "limits": ["Conversion is refused while vanilla connection still exists; use a unique tag to disconnect first.", "Exact NetworkName spelling/case is routing boundary.", "Matching server/client Foundation and actor enrollment required for remote authorization."],
        "deps": "BepInExPack; Core; Permissions; Persistence; Transactions 1.0.0"
    },
    {
        "name": "Runic Exploration", "version": "1.0.0",
        "purpose": "Client-side search and navigation for saved pins on map pixels the player has already explored.",
        "features": [
            "Searches saved known pins and optional already-replicated shared-map rows.",
            "Filters by category and source; groups exact one-metre/name/type display duplicates.",
            "Recognizes tombstone/bed/boss and explicit [boat], [cart], [animal]/[tame], [portal], [bed], [tombstone]/[grave] prefixes.",
            "Centers vanilla map after revalidation; shows LAST KNOWN distance, direction, elevation, and tombstone warning.",
            "When controlling a ship, shows LIVE LOCAL SHIP speed/sail, wind, biome, and selected-pin distance.",
            "Never reveals fog, scans entities, creates pins, tracks players, or confirms that an asset still exists.",
        ],
        "use": ["Open Valheim's normal large map; no separate activation hotkey.", "Click Search and type up to 48 characters; click Category and Source to cycle filters; click a result to center.", "Escape while search is focused releases text focus; closing/reopening map hides/restores panel state.", "Controller uses vanilla map controls; mouse/keyboard selects browser rows."],
        "verify": ["Panel appears on large map and remains available after typing/Escape.", "Unknown/unsaved/unexplored pins are absent; known saved pins appear as LAST KNOWN.", "Clicking a result centers map and produces distance/direction.", "Ship readout appears only while actually controlling a ship."],
        "limits": ["Client-only; no RPC, ownership, routefinding, live asset verification, or remote recovery.", "At most 24 displayed results; refine query when matches exceed visible rows.", "No-map worlds and too-small safe areas hide the panel."],
        "deps": "BepInExPack; Runic Core 1.0.0"
    },
    {
        "name": "Runic Awareness", "version": "1.0.0",
        "purpose": "Display-only explanations for already-known local food, effects, comfort, hovered contexts, and bounded building information.",
        "features": [
            "Shows active food names/timers and icon-bearing status effects.",
            "Shows comfort, shelter, Rested time, and captured vanilla comfort category winners.",
            "Captures item comparison from vanilla tooltip but yields while inventory is open.",
            "Shows bounded hovered production, agriculture, beehive, tameable, and building context.",
            "Optionally includes existing Runic Production status for locally owned reachable station.",
            "Each panel toggles independently; default anchor is MiddleLeft and controller gets readability scaling.",
        ],
        "use": ["No hotkey; panels appear automatically when their context is active.", "Configure panel toggles, anchor, scale, rows, and refresh interval.", "Hover supported stations/plants/beehives/tameables/build pieces within physical reach."],
        "verify": ["Food/effect/comfort text reflects current local state and timers.", "Hover supported objects and see only bounded context already available to client.", "Open inventory, map, build selector, chat, pause, text entry, trader, or popup; overlay yields completely.", "Remote Build Camera target does not reveal protected building detail."],
        "limits": ["No radar, hidden enemies, weather prediction, world scan, or gameplay mutation.", "Inventory comparison capture is compatibility-only and not drawn over inventory.", "Dedicated server is inert."],
        "deps": "BepInExPack; Runic Core 1.0.0"
    },
    {
        "name": "Runic Safety", "version": "1.0.0",
        "purpose": "Contextual confirmations, protected-item policy, death/recovery planning, migration backups, and compatibility admission.",
        "features": [
            "Requires identical repeat before removing occupied chest, ship/cart, or overwriting non-empty portal tag.",
            "Protects locked, equipped, quest, and configured rare items at audited destructive destinations.",
            "Rare incineration confirmation runs on client and authoritative owner.",
            "Preflights native grave recovery capacity and cooperates with Inventory topology/durable custody.",
            "Provides bounded checksummed migration backups with restore manifest before caller mutation.",
            "Required/Optional/Disabled compatibility gate compares game/module/topology/rules profiles.",
        ],
        "use": ["Repeat the same unchanged high-impact action within 4 seconds to confirm.", "Configure exact RarePrefabNames and confirmation toggles; admin bypass remains off unless deliberately required.", "Keep RemoteAdmissionPolicy=Required with matching profiles, or use Optional for first-run diagnosis."],
        "verify": ["First occupied-container or ship/cart removal warns; second identical action proceeds through vanilla.", "Locked/equipped/quest item is denied at protected destination.", "Configured rare item requires confirmation rather than disappearing immediately.", "Matching client joins under selected admission policy; mismatch gives an actionable denial."],
        "limits": ["Safety only declines/allows original actions; it does not authorize or perform them.", "No extra recovery chest or automatic multi-file restore.", "Unknown/incompatible evidence fails closed."],
        "deps": "BepInExPack; Core; Persistence; Transactions 1.0.0"
    },
    {
        "name": "Runic Sentinel", "version": "1.0.0",
        "purpose": "Bounded local plugin snapshot and RSA-signed policy review with optional/required pre-admission claims.",
        "features": [
            "Hashes loaded plugin files and canonicalizes identity/version/dependency evidence.",
            "Verifies strict RUNIC-SENTINEL/2 policy with pinned RSA-3072 public key and detached signature.",
            "Publishes bounded self-reported compatibility claim through Persistence handshake.",
            "Optional default records mismatch/monitor-only evidence without locking out first-run players.",
            "Required mode demands same valid signed policy; Disabled publishes local claim but installs no local evaluator.",
            "No telemetry, private key, kernel component, process scan, or world mutation.",
        ],
        "use": ["Leave Remote Admission.Policy=Optional until a signed policy/public-key pin is provisioned everywhere.", "For enforcement, install the same canonical policy/key pin on server and clients, switch server to Required, restart.", "Inspect bounded log/evidence only when diagnosing profile mismatch."],
        "verify": ["Optional first run reaches world even without signed policy and reports monitor-only state.", "A valid signed policy verifies and produces consistent digest/sequence/profile.", "If testing Required, a deliberately mismatched client is rejected before world admission, then matching client joins."],
        "limits": ["Client snapshot is honest self-report, not anti-cheat attestation.", "No bans or remote-admin elevation; Required can reject only current connection.", "Enabled=false is startup-inert."],
        "deps": "BepInExPack; Runic Core; Runic Persistence 1.0.0"
    },
    {
        "name": "Runic Velocity", "version": "1.0.0",
        "purpose": "Startup timing and integrity-checked local plugin-manifest cache for diagnostics.",
        "features": [
            "Records bounded milestones through first update, menu, network, world, and local-player readiness.",
            "Background-scans plugin DLL metadata/hashes without loading plugin code.",
            "Caches path, size, time, SHA-256, identity, version, dependencies, classification, and scan time.",
            "Reuses cache only when exact path/size/time/hash format validate; changed files are rehashed from one stable stream.",
            "Publishes startup.measure, startup.cache, and read-only status service.",
        ],
        "use": ["No controls or UI; settings are sampled once at startup.", "Change settings before next launch rather than during active scan.", "Use logs/status service when comparing cold/warm startup or plugin changes."],
        "verify": ["Launch reaches menu/world and timeline records expected readiness milestones.", "Second unchanged launch may reuse valid cache; modifying a plugin invalidates only affected evidence and rebuilds cleanly.", "No worker Unity errors or gameplay changes occur."],
        "limits": ["Not a preloader, dependency repair system, anti-cheat, Safe Start, or Server Forge.", "Local cache is a performance hint, never authorization.", "Enabled=false creates no timeline/worker/module/service."],
        "deps": "BepInExPack; Runic Core 1.0.0"
    },
    {
        "name": "Runic World Engine", "version": "1.0.0",
        "purpose": "Conservative read-only world-state observatory and bounded registry for Runic-owned ZDO data declarations.",
        "features": [
            "Measures aggregate ZDO population, peer count, create/destroy intervals, send/receive rates, and save/load duration.",
            "Hard maximum one aggregate sample per second; no ordinary-update world scan.",
            "Publishes zdo.observe and zdo.ownership typed services.",
            "Registry lets Runic modules declare prefab hashes, persistent/temporary keys, and schema versions.",
            "Rejects conflicting/unsafe declarations and preserves unknown third-party data.",
            "Optional periodic summaries are off by default and rate-limited."],
        "use": ["No player controls or UI.", "Enable Diagnostics.LogPeriodicSummary only when needed; interval is clamped 5–600 seconds.", "Gameplay modules register ownership declarations through Core, not direct user configuration."],
        "verify": ["World load reports observatory/ownership service ready.", "Optional summary shows bounded aggregate counts no faster than configured cadence.", "Normal play shows no ZDO deletion, rewriting, compaction, or forced sends."],
        "limits": ["Observe-only: no heat map, cleanup, quarantine, scheduling, or persistent keys of its own.", "Conflicts leave prior registry state unchanged.", "Same small runtime is safe on dedicated server."],
        "deps": "BepInExPack; Runic Core 1.0.0"
    },
]


def add_module_page(doc, index, module):
    title_block(doc, f"Module packet {index:02d} of {len(MODULES):02d}", f"{module['name']}  {module['version']}", module["purpose"])
    small_table(doc, ["Package role", "Dependencies"], [("Gameplay" if module["name"] not in {"Runic Core", "Runic Persistence", "Runic Permissions", "Runic Transactions", "Runic Sentinel", "Runic Velocity", "Runic World Engine"} else "Foundation / diagnostics", module["deps"])], [2100, 8066])
    heading(doc, "What it provides", 2)
    for item in module["features"]:
        bullet(doc, item)
    heading(doc, "How to use", 2)
    for item in module["use"]:
        bullet(doc, item)
    heading(doc, "Quick verification", 2)
    for item in module["verify"]:
        check(doc, item)
    heading(doc, "Expected limits / support notes", 2)
    for item in module["limits"]:
        bullet(doc, item)
    callout(doc, "PASS", "All checks above match, no item/world-state discrepancy is observed, and any denial includes a clear reason. If not, stop at the first exact failure and preserve its screenshot, reproduction steps, and session LogOutput.log.", PALE_GREEN)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_doc(doc)
    doc.core_properties.title = "Runic Mod Suite In-Game Testing Primer and Reference"
    doc.core_properties.subject = "Three-page test primer plus one-page reference packet for all 20 packaged Runic mods"
    doc.core_properties.author = "Chazman Mods"
    doc.core_properties.keywords = "Valheim, Runic, testing, controls, reference"
    add_primer_page_1(doc)
    page_break(doc)
    add_primer_page_2(doc)
    page_break(doc)
    add_primer_page_3(doc)
    for idx, module in enumerate(MODULES, 1):
        page_break(doc)
        add_module_page(doc, idx, module)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
