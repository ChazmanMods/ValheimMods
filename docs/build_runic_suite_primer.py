from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path(r"E:\Valheim Mods\ChazmanModsRepo\docs\Runic_Mod_Suite_Primer_and_Mini_Base_Walkthrough.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
BODY = RGBColor(36, 42, 48)
MUTED = RGBColor(92, 102, 112)
GOLD = RGBColor(181, 133, 31)
GREEN = RGBColor(48, 122, 74)
RED = RGBColor(155, 28, 28)
TURQUOISE = RGBColor(0, 128, 128)
WHITE = RGBColor(255, 255, 255)

PALE_BLUE = "E8EEF5"
PALE_GREY = "F2F4F7"
PALE_GOLD = "FFF4D6"
PALE_GREEN = "EAF4E8"
PALE_RED = "F9E9E7"
PALE_TURQUOISE = "E4F4F3"
NAVY = "0B2545"

bookmark_counter = 1
number_counter = 0
pending_page_break = False


def set_font(run, size=None, bold=None, color=None, italic=None, font="Calibri"):
    run.font.name = font
    rfonts = run._element.get_or_add_rPr().rFonts
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def fixed_table(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), "120")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[i]))
    # Mark the first row for assistive navigation.  This is also used by the
    # one-row callout tables so they are not exposed as headerless tables.
    if table.rows:
        repeat_header(table.rows[0])


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is not None:
        return
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def add_field(paragraph, field):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def add_bookmark(paragraph, name):
    global bookmark_counter
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_counter))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_counter))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)
    bookmark_counter += 1


def internal_link(paragraph, text, anchor, color=BLUE, bold=False):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), str(color))
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(c)
    rpr.append(u)
    if bold:
        rpr.append(OxmlElement("w:b"))
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, before, after, color in (
        ("Title", 28, 0, 8, INK),
        ("Subtitle", 14, 0, 10, MUTED),
        ("Heading 1", 16, 18, 10, BLUE),
        ("Heading 2", 13, 14, 7, BLUE),
        ("Heading 3", 12, 10, 5, DARK_BLUE),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    if "Compact" not in styles:
        compact = styles.add_style("Compact", WD_STYLE_TYPE.PARAGRAPH)
    else:
        compact = styles["Compact"]
    compact.font.name = "Calibri"
    compact.font.size = Pt(9)
    compact.font.color.rgb = BODY
    compact.paragraph_format.space_after = Pt(3)
    compact.paragraph_format.line_spacing = 1.08

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    set_font(hp.add_run("RUNIC MOD SUITE  |  PRIMER & MINI-BASE WALKTHROUGH"), 8, True, MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    set_font(fp.add_run("Chazman Mods  |  Page "), 8, False, MUTED)
    add_field(fp, "PAGE")

    props = doc.core_properties
    props.title = "Runic Mod Suite Primer and Mini-Base Walkthrough"
    props.subject = "A linked introduction and complete reference for the current Runic Valheim mod suite"
    props.author = "Chazman Mods"
    props.keywords = "Valheim, Runic, mod suite, primer, walkthrough, production, agriculture"


def page_break(doc):
    global pending_page_break
    pending_page_break = True


def title(doc, text, bookmark=None, kicker=None):
    global number_counter, pending_page_break
    number_counter = 0
    if kicker:
        p = doc.add_paragraph()
        if pending_page_break:
            p.paragraph_format.page_break_before = True
            pending_page_break = False
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(kicker.upper()), 8.5, True, GOLD)
    p = doc.add_paragraph(text, style="Heading 1")
    if pending_page_break:
        p.paragraph_format.page_break_before = True
        pending_page_break = False
    if bookmark:
        add_bookmark(p, bookmark)
    return p


def h2(doc, text, bookmark=None):
    global number_counter
    number_counter = 0
    p = doc.add_paragraph(text, style="Heading 2")
    if bookmark:
        add_bookmark(p, bookmark)
    return p


def h3(doc, text, bookmark=None):
    global number_counter
    number_counter = 0
    p = doc.add_paragraph(text, style="Heading 3")
    if bookmark:
        add_bookmark(p, bookmark)
    return p


def para(doc, text="", style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.56)
    p.add_run(text)
    return p


def numbered(doc, text):
    global number_counter
    number_counter += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(f"{number_counter}.  "), 10.5, True, DARK_BLUE)
    p.add_run(text)
    return p


def callout(doc, label, text, fill=PALE_GOLD, label_color=GOLD):
    table = doc.add_table(rows=1, cols=1)
    fixed_table(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(label.upper() + "  "), 9.5, True, label_color)
    set_font(p.add_run(text), 10.2, False, BODY)
    return table


def nav(doc, back="contents"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    internal_link(p, "Back to contents", back, DARK_BLUE)
    p.add_run("   •   ")
    internal_link(p, "Back to mini-base walkthrough", "walkthrough", DARK_BLUE)


def kv_table(doc, rows, widths=(2700, 6660), header=None):
    count = len(rows) + (1 if header else 0)
    table = doc.add_table(rows=count, cols=2)
    table.style = "Table Grid"
    fixed_table(table, list(widths))
    offset = 0
    if header:
        for i, value in enumerate(header):
            shade(table.cell(0, i), PALE_BLUE)
            p = table.cell(0, i).paragraphs[0]
            set_font(p.add_run(value), 9.5, True, INK)
        repeat_header(table.rows[0])
        offset = 1
    for r, (left, right) in enumerate(rows, start=offset):
        if r % 2 == offset % 2:
            shade(table.cell(r, 0), PALE_GREY)
            shade(table.cell(r, 1), PALE_GREY)
        p = table.cell(r, 0).paragraphs[0]
        set_font(p.add_run(left), 9.3, True, DARK_BLUE)
        p = table.cell(r, 1).paragraphs[0]
        set_font(p.add_run(right), 9.3, False, BODY)
    return table


def module_header(doc, name, version, tagline, bookmark):
    page_break(doc)
    title(doc, name, bookmark, "Detailed mod reference")
    p = doc.add_paragraph()
    set_font(p.add_run(version + "  |  "), 9.5, True, GOLD)
    set_font(p.add_run(tagline), 11.5, False, MUTED, True)
    nav(doc)


def control_table(doc, rows):
    return kv_table(doc, rows, widths=(3000, 6360), header=("Control", "Result"))


def content_link_paragraph(doc, label, anchor, description):
    p = doc.add_paragraph()
    internal_link(p, label, anchor, DARK_BLUE, True)
    p.add_run(" — " + description)
    return p


doc = Document()
configure(doc)

# Cover
p = doc.add_paragraph()
add_bookmark(p, "top")
p.paragraph_format.space_after = Pt(26)
set_font(p.add_run("CHAZMAN MODS"), 10, True, GOLD)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(60)
p.paragraph_format.space_after = Pt(8)
set_font(p.add_run("RUNIC"), 42, True, INK)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(18)
set_font(p.add_run("MOD SUITE PRIMER"), 24, True, BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(28)
set_font(p.add_run("Build a working mini base, then learn every mod in detail"), 15, False, MUTED, True)

table = doc.add_table(rows=3, cols=1)
fixed_table(table, [9360])
for i, (label, value, fill) in enumerate((
    ("WALKTHROUGH", "One starter base using every appropriate Runic system", PALE_BLUE),
    ("REFERENCE", "At least one dedicated page for each of the 16 current mods", PALE_GREY),
    ("NAVIGATION", "Linked key terms jump directly to definitions and exact instructions", PALE_GOLD),
)):
    cell = table.cell(i, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    set_font(p.add_run(label + "  "), 10, True, DARK_BLUE)
    set_font(p.add_run(value), 11, False, BODY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(34)
set_font(p.add_run("Current packaged suite • August 2026"), 10, True, GOLD)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(0)
set_font(p.add_run("Valheim + BepInEx 5 • Standalone Runic packages"), 9.5, False, MUTED)

# Contents
page_break(doc)
title(doc, "Contents", "contents", "Navigate the primer")
callout(doc, "How to use the links", "Blue underlined terms are internal links. Select one to jump to the detailed definition or exact control instructions. Each mod page links back here and to the walkthrough.", PALE_BLUE, BLUE)
h2(doc, "Part I — Build the mini base")
for label, anchor, desc in (
    ("1. Prepare the profile", "phase_prepare", "Install, verify, and understand ownership."),
    ("2. Plan the site", "phase_plan", "Lay out a compact base and its functional zones."),
    ("3. Build accurately", "phase_build", "Use detached camera and local-axis placement."),
    ("4. Organize gear and storage", "phase_storage", "Create safe rows, searchable stores, and display stands."),
    ("5. Plant the garden", "phase_garden", "Configure a grid, source seeds, preview, plant, and harvest."),
    ("6. Connect production", "phase_production", "Build input, fuel, output, relay, and replenishment links."),
    ("7. Add travel and navigation", "phase_travel", "Create portal routes and use last-known map search."),
    ("8. Daily operating loop", "daily_loop", "A short routine that exercises the whole suite."),
):
    content_link_paragraph(doc, label, anchor, desc)

h2(doc, "Part II — Detailed mod reference")
mods_for_contents = (
    ("Runic Agriculture", "mod_agriculture"), ("Runic Awareness", "mod_awareness"),
    ("Runic Build Camera", "mod_buildcamera"), ("Runic Crafting", "mod_crafting"),
    ("Runic Display Stands", "mod_displaystands"), ("Runic Exploration", "mod_exploration"),
    ("Runic Interaction", "mod_interaction"), ("Runic Inventory", "mod_inventory"),
    ("Runic Portals", "mod_portals"), ("Runic Precision Build Tool", "mod_precision"),
    ("Runic Production", "mod_production"), ("Runic Safety", "mod_safety"),
    ("Runic Sentinel", "mod_sentinel"), ("Runic Storage", "mod_storage"),
    ("Runic Velocity", "mod_velocity"), ("Runic World Engine", "mod_worldengine"),
)
table = doc.add_table(rows=8, cols=2)
fixed_table(table, [4680, 4680])
for idx, (name, anchor) in enumerate(mods_for_contents):
    cell = table.cell(idx % 8, idx // 8)
    shade(cell, PALE_GREY if idx % 2 else PALE_BLUE)
    p = cell.paragraphs[0]
    internal_link(p, name, anchor, DARK_BLUE, True)

# How to use
page_break(doc)
title(doc, "How this primer works", "walkthrough", "Part I — Mini-base walkthrough")
para(doc, "The walkthrough builds one compact base in stages. Each stage names the Runic mods being used and links important terms to their full reference pages. The detailed reference then explains scope, controls, safeguards, configuration, and a quick verification for every current package.")
h2(doc, "The mini-base goal")
bullet(doc, "A 10 m × 12 m workshop shell with a clear central aisle.")
bullet(doc, "A workstation wall, searchable storage bank, gear/display corner, and small protected garden.")
bullet(doc, "A real relay chain: mead ketill → relay chest → fermenter → finished-mead chest.")
bullet(doc, "An input/fuel/output production bay, a portal alcove, and an exploration board.")
bullet(doc, "Background integrity, startup, performance, safety, and awareness support.")
h2(doc, "Three rules that prevent confusion")
numbered(doc, "Use the same current package build on every process that needs the mod. Sentinel policy enforcement is a separate, deliberate administrative choice.")
numbered(doc, "Runic systems do not bypass Valheim ownership, wards, station access, placement validity, or ordinary resource costs.")
numbered(doc, "When a feature does not apply to a station or screen, the mod yields to vanilla behavior instead of inventing a substitute workflow.")
callout(doc, "Terminology", "“Local owner” means the game process currently authoritative for the object. “Eligible chest” means a loaded, in-range container that passes native access and the mod’s bounded query rules.", PALE_GOLD, GOLD)

# Base map
page_break(doc)
title(doc, "Mini-base floor plan", "base_plan", "Part I — Mini-base walkthrough")
para(doc, "Keep the layout small enough that the default 20–30 m nearby-container ranges overlap. The map is conceptual: rotate or mirror it to suit the terrain.")
table = doc.add_table(rows=4, cols=3)
table.style = "Table Grid"
fixed_table(table, [3120, 3120, 3120])
zones = [
    ("PORTAL ALCOVE", "Portal frame\nMap/directory access", PALE_BLUE),
    ("STORAGE BANK", "Labeled ingredient chests\nSearch + quick stack", PALE_GREY),
    ("GEAR CORNER", "Armor/item stands\nInventory staging", PALE_GOLD),
    ("GARDEN", "Cultivated 10×10 grid\nSeed reserve chest", PALE_GREEN),
    ("CENTRAL AISLE", "Open sight lines\nAwareness + interaction", "FFFFFF"),
    ("CRAFT WALL", "Workbench / forge\nNearby materials", PALE_BLUE),
    ("KETILL", "Recipe input\nReplenishment exemplar", PALE_TURQUOISE),
    ("RELAY CHEST", "Ketill output\nFermenter input", PALE_GOLD),
    ("FERMENTER", "Base input\nFinished-mead output", PALE_TURQUOISE),
    ("SMELTER / OVEN", "Input + fuel input\nOutput chests", PALE_RED),
    ("REPLENISHMENT", "Target item exemplar\nReserve quantity", PALE_TURQUOISE),
    ("EXPANSION", "Lamp/fire fuel\nFuture stations", PALE_GREY),
]
for i, (label, text_value, fill) in enumerate(zones):
    cell = table.cell(i // 3, i % 3)
    shade(cell, fill)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(label + "\n"), 9.5, True, DARK_BLUE)
    set_font(p.add_run(text_value), 9, False, BODY)
h2(doc, "Recommended chest labels")
kv_table(doc, [
    ("Raw materials", "Wood, stone, ore, coal, resin, crops, seeds, and recipe ingredients."),
    ("Relay", "Intermediate product shared between an upstream output and downstream input."),
    ("Finished goods", "Bars, cooked food, mead, or other completed station output."),
    ("Replenishment", "Contains a physical exemplar of the item to keep at reserve quantity."),
])

# Phase 1
page_break(doc)
title(doc, "1. Prepare the profile", "phase_prepare", "Mini-base phase")
para(doc, "Start with a clean, current mod profile and verify the quiet background services before building anything.")
h2(doc, "Install and first launch")
numbered(doc, "Install the 16 current Runic packages and their declared BepInEx dependency. Do not mix retired Foundation packages into this primer’s setup.")
numbered(doc, "Launch once. Confirm BepInEx loads each package without an exception. Runic Velocity may report startup milestones; Runic World Engine normally remains quiet unless summaries are enabled.")
numbered(doc, "If using Runic Sentinel, choose Disabled, Optional, or Required deliberately. Required mode needs the same valid signed policy and pinned public key on server and clients.")
numbered(doc, "Open Configuration Manager and review ranges, keybinds, and UI settings. Runtime-applicable settings take effect without rebuilding the package; startup-only diagnostics require a restart.")
h2(doc, "Mods active in this phase")
p = doc.add_paragraph()
internal_link(p, "Runic Velocity", "mod_velocity", DARK_BLUE, True)
p.add_run(" shortens avoidable startup work through a verified local manifest cache. ")
internal_link(p, "Runic World Engine", "mod_worldengine", DARK_BLUE, True)
p.add_run(" measures aggregate world/network activity. ")
internal_link(p, "Runic Sentinel", "mod_sentinel", DARK_BLUE, True)
p.add_run(" can validate a signed mod policy. ")
internal_link(p, "Runic Safety", "mod_safety", DARK_BLUE, True)
p.add_run(" guards high-impact actions later in the build.")
callout(doc, "Expected result", "The world loads normally. Diagnostic mods do not add a gameplay menu, and monitor-only modes do not block entry.", PALE_GREEN, GREEN)

# Phase 2 plan
page_break(doc)
title(doc, "2. Plan the site", "phase_plan", "Mini-base phase")
para(doc, "Choose flat ground near water and leave enough room for a cultivated plot. Put the storage bank near the center so crafting, agriculture, and replenishment queries can reach it without extreme ranges.")
h2(doc, "Mark the zones")
numbered(doc, "Place temporary beams for the 10 m × 12 m shell, central aisle, garden edge, portal alcove, and production wall.")
numbered(doc, "Put the ingredient storage bank between the garden and production line. This reduces travel and keeps bounded chest searches predictable.")
numbered(doc, "Leave station interaction points and output bays unobstructed: smelter front bay, oven doors, fermenter spout, and fuel controls must remain targetable.")
numbered(doc, "Reserve a chest between the mead ketill and fermenter. It will be both an upstream output and downstream input.")
h2(doc, "Use Awareness while planning")
p = doc.add_paragraph()
p.add_run("The ")
internal_link(p, "context panels", "awareness_panels", DARK_BLUE, True)
p.add_run(" can show building position/rotation/support and comfort information without revealing hidden world data. Panels yield when inventory, map, chat, pause, or build menus need the screen.")
callout(doc, "Range discipline", "Keep ordinary material chests inside Runic Crafting’s default 20 m cap and seed sources inside Runic Agriculture’s configured range. Increase ranges only when the physical layout genuinely needs it.", PALE_GOLD, GOLD)

# Phase 3 build
page_break(doc)
title(doc, "3. Build accurately", "phase_build", "Mini-base phase")
para(doc, "Use the detached camera for reach and overview, then use local-axis precision controls for pieces whose pitch, roll, or offset must follow the object itself.")
h2(doc, "Build sequence")
numbered(doc, "Equip the hammer and open build mode. Press B to toggle the detached Build Camera; move and look normally, use jump/crouch vertically, and run for faster movement.")
numbered(doc, "Press P only while a hammer piece preview is active to enable Precision Build. The key does nothing outside that context.")
numbered(doc, "Place the floor and posts. Rotate the selected piece around its local yaw, pitch, or roll; a sloped beam therefore pitches relative to itself, not only the world axes.")
numbered(doc, "Use the precision search/favorites/recent controls to return to repeated structural pieces. Use guides and numeric modes for offsets and snaps.")
numbered(doc, "Finish the shell, production wall, and unobstructed station bays. Leave the garden open to the sky where the crop requires it.")
h2(doc, "Linked instructions")
p = doc.add_paragraph()
internal_link(p, "Detached camera limits", "buildcamera_limits", DARK_BLUE, True)
p.add_run(" explain range, pickup, vulnerability, and compatibility. ")
internal_link(p, "Precision controls", "precision_controls", DARK_BLUE, True)
p.add_run(" list every keyboard/mouse command. ")
internal_link(p, "Safety confirmations", "safety_confirmations", DARK_BLUE, True)
p.add_run(" explain repeat-to-confirm protection for destructive actions.")
callout(doc, "Expected result", "You can inspect the entire frame without moving the avatar, while every placed piece still uses vanilla cost, station, ward, and placement validation.", PALE_GREEN, GREEN)

# Phase 4 storage
page_break(doc)
title(doc, "4. Organize gear and storage", "phase_storage", "Mini-base phase")
h2(doc, "Prepare the player inventory")
numbered(doc, "Open inventory and identify the labeled bottom row: Helmet, Chest, Legs, Cape, Utility, Quick 1, Quick 2, and Quick 3.")
numbered(doc, "Equip gear into its role. Put food, mead, or a utility item in Quick 1–3 and use Alt+1, Alt+2, or Alt+3.")
numbered(doc, "Alt+right-click any player slot to lock it. The yellow outline is the confirmation; no HUD message is required. Alt+L is the optional focused-slot fallback.")
h2(doc, "Build the storage bank")
numbered(doc, "Create separate chests for raw materials, seeds/crops, fuel, recipe ingredients, relay items, finished output, and replenishment targets.")
numbered(doc, "Use Alt+Q to quick-stack matching carried items. Open a chest and use Alt+A to store all eligible carried items, Alt+S to sort that chest, or Alt+R to restock configured carried targets.")
numbered(doc, "Use Alt+F to open nearby-item search. Type or select an item; every nearby matching chest receives a bright yellow highlight for 15 seconds.")
h2(doc, "Finish the gear corner")
numbered(doc, "Place an armor stand and item stand. Hold Left Alt while interacting to access the compact native container workflow or take one item.")
numbered(doc, "Use Runic Interaction’s Alt+click stack transfer in an open inventory. Holding Use repeats supported station fuel/input at native cadence.")
p = doc.add_paragraph()
p.add_run("Details: ")
internal_link(p, "protected inventory row", "inventory_roles", DARK_BLUE, True)
p.add_run(" • ")
internal_link(p, "storage controls", "storage_controls", DARK_BLUE, True)
p.add_run(" • ")
internal_link(p, "search menu", "storage_search", DARK_BLUE, True)
p.add_run(" • ")
internal_link(p, "display stand rules", "display_rules", DARK_BLUE, True)

# Garden
page_break(doc)
title(doc, "5. Plant the garden", "phase_garden", "Mini-base phase")
para(doc, "Runic Agriculture plants the exact player-selected pattern, limited by available resources and valid planting positions—not by a separate carried-seed budget.")
h2(doc, "Make a 10 × 10 starter grid")
numbered(doc, "Cultivate an area large enough for the crop’s minimum spacing. Equip the cultivator and select the seed recipe.")
numbered(doc, "Choose Grid. Use Alt+wheel for rows and Shift+wheel for columns until the preview is 10 × 10. Rows and columns can each be 1–40, with a hard maximum of 1,600 live preview cells.")
numbered(doc, "Use Alt+Shift+wheel for spacing. The mod will not allow spacing below the crop’s own minimum. Use the bare wheel to rotate the whole pattern.")
numbered(doc, "Check the preview colors. Green is ready, red is a resource shortfall, and amber is an invalid position such as terrain, spacing, access, or biome failure.")
numbered(doc, "Left-click once to plant every ready cell in the visible configuration. Resources are drawn from personal inventory first, then eligible nearby chests.")
h2(doc, "Shortages and fill order")
p = doc.add_paragraph()
p.add_run("For ")
internal_link(p, "Grid and Row", "agriculture_fill", DARK_BLUE, True)
p.add_run(", the resource-limited fill proceeds from the player’s right to left, then front to back. Shaped patterns fill from the center outward. Red ghosts appear only where requirements are unavailable; invalid positions are amber.")
h2(doc, "Harvest")
para(doc, "Use Alt+E for bounded area harvest. Alt+T offers replanting where supported. The bottom native BuildHints bar shows the active shape, dimensions, spacing, rotation, and ready count.")
callout(doc, "Expected result", "A single left-click plants the selected 10 × 10 configuration wherever terrain and resources allow. A 1 × 1 configuration is the deliberate single-plant mode.", PALE_GREEN, GREEN)

# Production overview
page_break(doc)
title(doc, "6. Connect production", "phase_production", "Mini-base phase")
para(doc, "Runic Production uses a simple two-click link: choose a role on the station, then repeat the same gesture on the target chest within 30 seconds. A chest may serve multiple roles, including output for one station and input for the next.")
h2(doc, "Role gestures")
control_table(doc, [
    ("Alt+Left Mouse", "Input. Aim at a station, then repeat on an input chest."),
    ("Alt+Left Mouse on native fuel control", "Fuel Input. The message explicitly says Fuel Input rather than ordinary Input."),
    ("Alt+Right Mouse", "Output. Aim at the station or its native output location, then repeat on an output chest."),
    ("Alt+Middle Mouse", "Replenishment. Repeat on each replenishment chest."),
    ("Shift+Alt + the same button", "Removal mode. Use the same shifted gesture on station and linked chest."),
])
h2(doc, "Visual confirmation")
kv_table(doc, [
    ("Green ring", "Input or Fuel Input chest."),
    ("Yellow ring", "Output chest."),
    ("Turquoise ring", "Replenishment chest; multiple are allowed."),
])
p = doc.add_paragraph()
p.add_run("When the station is focused, its hover panel lists ")
internal_link(p, "existing production links", "production_links", DARK_BLUE, True)
p.add_run(" and empty roles. Pointing at a station input/output region also reveals the linked chest rings.")
callout(doc, "Native authority", "Links and automation operate only when station, chest, ownership, access, range, and native state can be proven. An indeterminate station may pause for that session rather than risk duplicating or losing items.", PALE_GOLD, GOLD)

# Production chain
page_break(doc)
title(doc, "Build the ketill-to-fermenter line", "production_chain", "Mini-base phase 6")
h2(doc, "A. Mead ketill")
numbered(doc, "Link one or more ingredient chests as Input with Alt+Left Mouse.")
numbered(doc, "Link the relay chest as Output with Alt+Right Mouse.")
numbered(doc, "Link a replenishment chest with Alt+Middle Mouse. Put one physical exemplar of the desired mead base in it and set/retain the reserve quantity. The exemplar tells the recipe station what to make.")
h2(doc, "B. Relay chest")
para(doc, "Use the same physical relay chest as the ketill’s Output and the fermenter’s Input. This is supported; the chest is not limited to one role or one station.")
h2(doc, "C. Fermenter")
numbered(doc, "Focus the fermenter and link the relay chest as Input with Alt+Left Mouse.")
numbered(doc, "Link one or more finished-mead chests as Output with Alt+Right Mouse. Matching existing mead stacks help route the produced type.")
numbered(doc, "Optionally link replenishment chests with Alt+Middle Mouse for reserve-driven downstream production.")
h2(doc, "D. Smelter, oven, fires, and lamps")
bullet(doc, "Smelter: input ore chest, fuel-input coal chest, output bar chest.")
bullet(doc, "Oven/cooking station: input food, fuel input where the station has native fuel, output cooked food, optional replenishment.")
bullet(doc, "Fire or lamp: designate a Fuel Input chest by starting on its native fuel control.")
bullet(doc, "A Replenishment chest pulls ingredients from eligible nearby chests to restore the exemplar item when it falls below reserve.")
callout(doc, "Important", "An Output link alone does not guess a recipe. Recipe stations need an exemplar in a linked Replenishment chest to identify the desired product.", PALE_RED, RED)

# Travel
page_break(doc)
title(doc, "7. Add travel and navigation", "phase_travel", "Mini-base phase")
h2(doc, "Create the portal alcove")
numbered(doc, "Place a standard paired portal or aim at a portal and use Shift+E to open Runic editing.")
numbered(doc, "Choose Public, Private, or Group network mode; enter the case-sensitive network name and select Both, Arrive, or Depart endpoint behavior.")
numbered(doc, "Walk into a Depart/Both network portal to open the large-map destination picker. Select an eligible arrival endpoint or press Esc to cancel.")
numbered(doc, "Press P on the large map to view the portal directory. Standard Pair remains ordinary Valheim tag matching.")
h2(doc, "Use the map as an operations board")
numbered(doc, "Create saved pins such as [portal] Main Base, [boat] Longship, [cart] Copper, [animal] Boars, [bed] Outpost, and [tombstone] Recovery.")
numbered(doc, "Open the large map and use Runic Exploration search/filter controls. Results are saved pins in already explored map pixels and are explicitly last-known, not live tracking.")
numbered(doc, "Select a pin to see straight-line distance, direction, and elevation. A controlled ship receives a separate live-local readout.")
p = doc.add_paragraph()
p.add_run("Reference: ")
internal_link(p, "portal modes and commands", "portal_modes", DARK_BLUE, True)
p.add_run(" • ")
internal_link(p, "exploration tags and limits", "exploration_tags", DARK_BLUE, True)
callout(doc, "Expected result", "The portal network respects its permission mode, while exploration tools search only information the player has already saved and explored.", PALE_GREEN, GREEN)

# Daily loop
page_break(doc)
title(doc, "8. Daily operating loop", "daily_loop", "Mini-base phase")
numbered(doc, "Return through the portal and use the map directory or a saved [portal] pin to confirm the route.")
numbered(doc, "Open inventory: equip role items, use Quick slots, and verify any locked-slot outlines.")
numbered(doc, "Press Alt+F and select the material you need. Follow the yellow chest highlight; Alt+Q quick-stacks carried matches.")
numbered(doc, "Check production station hover panels and link rings. Refill exemplar/reserve chests only when changing the desired replenishment product.")
numbered(doc, "Craft/build normally; nearby eligible chests satisfy deficits. One repair-button press repairs every currently repairable worn item when Repair All is enabled.")
numbered(doc, "Equip the cultivator, choose dimensions/spacing, inspect ghost colors, and left-click to plant. Use Alt+E to harvest when ready.")
numbered(doc, "Use Build Camera and Precision Build for expansion. Repeat protected destructive actions only when Runic Safety’s confirmation is intentional.")
numbered(doc, "If diagnosing performance or policy, consult logs from World Engine, Velocity, or Sentinel; they do not create a normal gameplay menu.")
h2(doc, "The suite at a glance")
kv_table(doc, [
    ("Build", "Build Camera + Precision Build Tool + Crafting + Safety"),
    ("Organize", "Inventory + Storage + Display Stands + Interaction"),
    ("Grow / make", "Agriculture + Production + Crafting"),
    ("Travel / know", "Portals + Exploration + Awareness"),
    ("Operate safely", "Sentinel + Velocity + World Engine + Safety"),
])
callout(doc, "Next", "Use Part II whenever a linked concept, control, limit, or troubleshooting boundary needs a precise answer.", PALE_BLUE, BLUE)

# Detailed modules
module_header(doc, "Runic Agriculture", "1.0.0", "Pattern planting and bounded area harvest for the local cultivator.", "mod_agriculture")
h2(doc, "What it does")
para(doc, "Runic Agriculture replaces one-at-a-time cultivator placement with a native bottom-screen control bar and a live multi-cell preview. The player chooses the shape, rows, columns, spacing, rotation, and mirror/taper options. One left-click plants every ready cell. The mod consumes all normal piece requirements from personal inventory first, then eligible nearby chests.")
h2(doc, "Shapes, size, and resource behavior", "agriculture_fill")
bullet(doc, "Shapes: Row, Grid, Circle, Star, Right Triangle, Half Circle, and Trapezoid. Grid is the default.")
bullet(doc, "Rows and columns are independently adjustable from 1–40. The live pattern hard-limit is 1,600 cells.")
bullet(doc, "Grid/Row resource shortages fill from the player’s right to left, then front to back. Shaped patterns fill center-out.")
bullet(doc, "Spacing never goes below the selected crop’s native minimum. Seed count is not artificially limited to a carried-only budget.")
h2(doc, "Ghost colors", "agriculture_ghosts")
kv_table(doc, [("Green", "Ready and resource-backed."), ("Red", "Requirements unavailable after personal inventory and eligible nearby chests are counted."), ("Amber", "Invalid terrain, crop spacing, biome, access, or other native placement rule.")])
callout(doc, "No blue resource state", "Blue ghosts are not the resource-shortage signal. If a blue-looking preview appears, capture the current package version, selected crop, and log because current behavior defines red for missing requirements.", PALE_BLUE, BLUE)

page_break(doc)
title(doc, "Runic Agriculture — controls and setup", "agriculture_controls", "Detailed mod reference")
nav(doc)
control_table(doc, [
    ("Left Mouse", "Plant the entire ready portion of the visible configuration."),
    ("Alt+Build Menu or Alt+O", "Cycle pattern shape; Numpad 1–7 selects directly."),
    ("Mouse Wheel", "Rotate the pattern around its vertical/Z axis."),
    ("Alt+Wheel", "Change rows."), ("Shift+Wheel", "Change columns."),
    ("Alt+Shift+Wheel", "Change spacing, clamped to crop minimum."),
    ("Alt+Shift+L", "Mirror the pattern."),
    ("Alt+E", "Bounded area harvest."), ("Alt+T", "Replant offer where supported."),
])
h2(doc, "Mini-base setup")
numbered(doc, "Put seed chests within the configured 1–30 m range; default is 30 m.")
numbered(doc, "Select Grid and make a 10 × 10 layout. Reduce to 1 × 1 only when a single plant is desired.")
numbered(doc, "Rotate and space until the desired cells are green. Left-click once.")
h2(doc, "Boundaries")
para(doc, "The mod is local-owner and cultivator focused. It does not bypass native resource cost, cultivated-ground, biome, spacing, ward, or placement checks. Invalid cells are skipped rather than forced.")

module_header(doc, "Runic Awareness", "1.0.0", "Context-sensitive local information panels with no hidden-data scan.", "mod_awareness")
h2(doc, "What it does", "awareness_panels")
para(doc, "Runic Awareness displays information the local client already knows: current food and effects, comfort contributors, hovered item comparisons, and build-piece details such as position, rotation, and support. It is an interpretation layer, not a radar or world scanner.")
h2(doc, "Where it appears")
bullet(doc, "Default placement is Middle Left; individual panels, scale, and visibility are configurable.")
bullet(doc, "Panels appear only when their context is relevant and yield to inventory, map, build menu, chat, pause, and other blocking interfaces.")
bullet(doc, "There is no gameplay hotkey required. Information updates from the current local view.")
h2(doc, "Mini-base uses")
bullet(doc, "Check comfort contributors while placing the bed, fire, and shelter."),
bullet(doc, "Read support/rotation data while building the shell."),
bullet(doc, "Compare an item under the crosshair with currently equipped gear."),
h2(doc, "Boundaries")
para(doc, "Awareness does not reveal fog, hidden entities, remote inventories, private state, or unobserved world objects. It sends no gameplay RPC and makes no world mutation.")
callout(doc, "Verify", "Walk from outdoors into the finished shelter, then aim at a build piece. Comfort and build sections should change contextually and disappear when the inventory or map takes focus.", PALE_GREEN, GREEN)

module_header(doc, "Runic Build Camera", "1.0.0", "A detached construction camera that preserves normal build rules.", "mod_buildcamera")
h2(doc, "What it does")
para(doc, "Runic Build Camera detaches the build view from the player while the hammer is in use. It lets the camera fly around a structure, place within configured range, optionally pick up loose world drops, and keep a Wisplight following the active view.")
h2(doc, "Controls and limits", "buildcamera_limits")
control_table(doc, [
    ("B", "Toggle detached build camera."),
    ("Normal move/look", "Move the detached view."),
    ("Jump / Crouch", "Move camera vertically."),
    ("Run", "Move the camera faster."),
])
bullet(doc, "Default maximum camera distance is 60 m; hard cap 100 m. Action distance is also hard-capped at 100 m, with a bounded ray distance of 50 m.")
bullet(doc, "Optional pickup covers loose world drops only—not containers.")
bullet(doc, "The avatar remains in the world and can still take damage.")
h2(doc, "Compatibility and rules")
para(doc, "Every action still uses ordinary costs, stations, wards, placement validation, and local authority. Runic Build Camera is compatible with Runic Precision Build Tool. It intentionally guards against the conflicting Build Camera Custom Hammers Edition plugin.")
callout(doc, "Verify", "Toggle B in hammer build mode, fly to the roof edge, and place one valid piece. Confirm the player body remains at the original spot and the build consumes ordinary materials.", PALE_GREEN, GREEN)

module_header(doc, "Runic Crafting", "1.0.0", "Nearby-container crafting/building, workshop policy, and Repair All.", "mod_crafting")
h2(doc, "What it does")
para(doc, "Runic Crafting lets normal recipes and permitted building pieces satisfy missing requirements from eligible nearby containers. It consumes carried materials first, then plans a bounded chest withdrawal with local authority, validation, and rollback. Requirement UI reports carried, nearby, total, and missing counts.")
h2(doc, "Normal workflow")
bullet(doc, "Craft or build with the ordinary Valheim button; there is no extra material-use hotkey."),
bullet(doc, "The active station range is capped by Materials.RangeCapMeters (default 20 m, configurable 1–50 m)."),
bullet(doc, "Personal/private container types are excluded by default; native ward/access checks always apply."),
bullet(doc, "Stationless pieces use an explicit allow/deny list and a player-centered range. Invalid rules fail closed."),
h2(doc, "Repair All")
para(doc, "When enabled, one press of the station repair button repairs every worn item that vanilla currently considers repairable. Protected role items from Runic Inventory are temporarily admitted to the repair check without removing their protection.")
h2(doc, "Workshop Access")
para(doc, "Station use and local-material consumption are separate policies. Owners can inspect or change them with F5 commands such as runiccrafting_access show, station <policy>, materials <policy>, approve, unapprove, and group.")
callout(doc, "Verify", "Put one recipe ingredient in inventory and the remainder in an eligible chest. The UI should show the combined total; crafting should consume carried items first and the exact chest deficit second.", PALE_GREEN, GREEN)

module_header(doc, "Runic Display Stands", "1.3.1", "Native container workflows for item and armor stands.", "mod_displaystands")
h2(doc, "What it does", "display_rules")
para(doc, "Runic Display Stands gives configured item and armor stands a compact native container interface. Item stands accept ordinary items. Armor stands organize one item per supported armor category and reject duplicates or invalid categories.")
h2(doc, "Interaction")
control_table(doc, [
    ("Hold Left Alt while interacting", "Access the stand workflow or take one item, depending on context."),
    ("Take all", "Return all stored stand items through the native inventory path."),
    ("Use/equip / armor switch", "Swap the stand loadout with equipped armor and the first matching hotbar weapon/shield where supported."),
])
h2(doc, "Safety and preservation")
bullet(doc, "Item metadata is preserved through the move."),
bullet(doc, "If an item cannot return safely, the operation is rejected or the item is dropped through a visible fallback rather than silently deleted."),
bullet(doc, "The host-synchronized configuration determines which prefabs receive stand behavior. All interacting players should install the same package."),
h2(doc, "Mini-base use")
para(doc, "Use the armor stand as a ready-loadout station beside the bed and portal. Use item stands for a weapon/tool display without treating them as general bulk storage.")
callout(doc, "Verify", "Place one valid item, take it back, then exercise an armor loadout swap. Confirm duplicates are rejected and item quality/durability remain intact.", PALE_GREEN, GREEN)

module_header(doc, "Runic Exploration", "1.0.0", "Search and navigation for already-known map pins.", "mod_exploration")
h2(doc, "What it does")
para(doc, "Runic Exploration adds client-side search, filters, distance, direction, and elevation to the large map. It indexes saved pins only when their pixels are already explored. Results are last-known notes, not a live entity tracker.")
h2(doc, "Recognized tags", "exploration_tags")
kv_table(doc, [
    ("[boat]", "Saved vessel location."), ("[cart]", "Saved cart location."),
    ("[animal] or [tame]", "Saved animal location."), ("[portal]", "Saved portal location."),
    ("[bed]", "Saved bed/outpost location."), ("[tombstone] or [grave]", "Recovery location with warning emphasis."),
])
h2(doc, "Interface and limits")
bullet(doc, "The search/filter panel appears on the large map and supports mouse/keyboard input."),
bullet(doc, "Controller map controls remain vanilla; the extension does not steal unrelated controller actions."),
bullet(doc, "Selecting a known pin reports straight-line distance, direction, and elevation."),
bullet(doc, "A ship readout is live-local only while the player is actually controlling that ship."),
h2(doc, "Boundaries")
para(doc, "The mod does not reveal fog, scan entities, locate unsaved objects, or send gameplay RPC. Update pins manually when a boat, cart, or animal moves.")
callout(doc, "Verify", "Create [boat] Shore and [tombstone] Test pins in explored terrain. Search each on the large map and confirm the result is labeled as last-known.", PALE_GREEN, GREEN)

module_header(doc, "Runic Interaction", "1.0.0", "Faster repeated use, stack transfer, and small quality-of-life safeguards.", "mod_interaction")
h2(doc, "What it does")
para(doc, "Runic Interaction improves ordinary actions without adding persistent world state. It repeats supported station fuel/input while Use is held, provides full-stack inventory transfer, optionally auto-closes doors, restores eligible session equipment/menu context, validates text entry, and can refuse configured world-item pickups.")
h2(doc, "Controls")
control_table(doc, [
    ("Hold Use", "Repeat supported native station fuel/input at the station’s normal cadence."),
    ("Alt+click in open inventory", "Transfer the full stack through the native inventory path."),
    ("Alt+Use on a denied world drop", "Bypass the local pickup filter for that intentional pickup."),
])
h2(doc, "Optional behavior")
bullet(doc, "Door auto-close is off by default and still checks obstruction and access."),
bullet(doc, "Pickup filters match exact prefab IDs or safe shared-name tokens; denied drops remain visible in the world."),
bullet(doc, "Equipment/menu restoration is session-scoped and yields when state cannot be proven."),
h2(doc, "Mini-base use")
para(doc, "Hold Use to feed supported stations manually during setup, Alt+click stacks between the player and a chest, and use the pickup filter to leave unwanted loose drops on the ground.")
callout(doc, "Verify", "Alt+click a test stack into a chest, then hold Use on a supported input. Confirm the stack moves once and repeat use follows native timing.", PALE_GREEN, GREEN)

module_header(doc, "Runic Inventory", "1.0.0", "A role-aware native bottom row, quick slots, locks, and safe sorting.", "mod_inventory")
h2(doc, "What it does", "inventory_roles")
para(doc, "Runic Inventory assigns the existing eight-cell bottom row to Helmet, Chest, Legs, Cape, Utility, Quick 1, Quick 2, and Quick 3. It does not add a second inventory or increase capacity. Role labels and outlines appear only while inventory is open, and dialogs render above them.")
h2(doc, "Controls")
control_table(doc, [
    ("Alt+1 / Alt+2 / Alt+3", "Use Quick 1, Quick 2, or Quick 3."),
    ("Alt+I", "Sort only configured safe general rows; the hotbar and special row are excluded."),
    ("Alt+right-click player slot", "Silently lock/unlock the slot; yellow outline means locked."),
    ("Alt+L", "Optional focused-slot lock fallback."),
])
h2(doc, "Role behavior")
bullet(doc, "Recognized gear can auto-equip into an empty role; replacement uses a lossless swap path."),
bullet(doc, "Role and locked slots are protected from unsafe sort/store/consolidation operations."),
bullet(doc, "The repair allowance lets a workstation repair protected equipped armor without stripping protection."),
bullet(doc, "Controller chords use validated existing Valheim actions and are configurable."),
h2(doc, "Mini-base use")
para(doc, "Keep a hammer/cultivator on the hotbar, armor in the first five role cells, and food/mead in Quick 1–3. Lock any irreplaceable carried item outside those roles.")
callout(doc, "Verify", "Move the belt out and back into Utility, lock a general slot, open Split Stack, then repair at a workstation. The role outline returns, the dialog stays above it, and eligible armor remains repairable.", PALE_GREEN, GREEN)

module_header(doc, "Runic Portals", "1.0.0", "Vanilla pairs plus permission-aware named portal networks.", "mod_portals")
h2(doc, "What it does", "portal_modes")
para(doc, "Runic Portals preserves Standard Pair behavior and adds named networks with Public, Private, or Group permissions. Network names are case-sensitive. Each endpoint can allow Both directions, Arrive only, or Depart only.")
h2(doc, "Controls")
control_table(doc, [
    ("E", "Normal vanilla portal/tag interaction."),
    ("Shift+E while aiming at portal", "Open Runic portal editor."),
    ("Walk into Depart/Both network endpoint", "Open destination picker on the large map."),
    ("P on large map / Esc in picker", "Open portal directory / cancel destination selection."),
])
h2(doc, "Editor command forms")
para(doc, "Examples include network|public|NETWORK|NAME|both, corresponding private/group forms, and standard to restore vanilla pairing. Replacing or restoring a route uses a repeat confirmation where required.")
h2(doc, "Permissions and ownership")
bullet(doc, "Public allows eligible players; Private and Group enforce the configured identity/group rules."),
bullet(doc, "Group chat commands integrate with the group policy workflow."),
bullet(doc, "Portal changes require the native local owner and do not seize remote ownership."),
callout(doc, "Verify", "Create two matching network endpoints, one Depart and one Arrive. Confirm an allowed player can choose the destination and a denied player receives a permission denial.", PALE_GREEN, GREEN)

module_header(doc, "Runic Precision Build Tool", "2.0.1", "Local-axis placement, offsets, guides, search, favorites, and history.", "mod_precision")
h2(doc, "Activation and local axes")
para(doc, "P toggles Precision Build only when the hammer is in build mode with an active piece preview. Yaw, pitch, and roll are applied relative to the selected object, so an already-rotated beam can pitch around its own local X axis.")
h2(doc, "Core controls", "precision_controls")
control_table(doc, [
    ("P", "Toggle precision mode in hammer build preview only."),
    ("Wheel / Alt+Wheel", "Local yaw / local pitch."),
    ("Shift+Wheel / Shift+Alt+Wheel", "Local roll."),
    ("V", "Fine adjustment modifier."),
    ("Alt+Left/Right", "Sway; local lateral offset."),
    ("Alt+Up/Down", "Heave; local vertical offset."),
    ("Alt+PageUp/PageDown", "Surge; local forward/back offset."),
    ("G", "Toggle guides."), ("F10", "Reset current precision transform."),
])
h2(doc, "Boundaries")
para(doc, "Precision controls alter the preview transform, not validation. Ordinary costs, support, station, ward, range, and collision rules still decide whether placement succeeds. Keyboard/mouse is the promised control surface; controller parity is not implied.")

page_break(doc)
title(doc, "Runic Precision Build Tool — advanced controls", "precision_advanced", "Detailed mod reference")
nav(doc)
h2(doc, "Numeric modes")
control_table(doc, [
    ("Numpad 0", "Rotation mode."), ("Numpad 1 / 2 / 3", "Select local axes."),
    ("Numpad 4 / 5 / 6", "Select X/Y/Z adjustment."), ("Numpad 7", "Position mode."),
    ("Numpad 8", "Full transform mode."), ("Numpad 9", "Snap mode."),
    ("Numpad .", "Repeat last adjustment."), ("Shift+Numpad key", "Reset that mode/axis."),
])
h2(doc, "Build-menu productivity")
control_table(doc, [
    ("F6", "Search build catalog; cycle all when query is empty."),
    ("F7", "Toggle favorite."), ("F8", "Cycle favorites."),
    ("F9", "Cycle recent pieces."),
    ("F11 with precision active and menu closed", "Undo supported precision placement."),
    ("F12 with precision active and menu closed", "Repair supported target."),
])

module_header(doc, "Runic Production", "1.0.0", "Multi-chest station links, relay chains, fuel routing, and reserve replenishment.", "mod_production")
h2(doc, "What it does", "production_links")
para(doc, "Runic Production links loaded, accessible chests to compatible station roles. Multiple Input, Fuel Input, Output, and Replenishment chests are supported where the station exposes those functions. The same chest may hold multiple roles and connect multiple stations, enabling true production lines.")
h2(doc, "Two-click link workflow")
numbered(doc, "Aim at the station—or its native fuel/output interaction region—and press the role gesture.")
numbered(doc, "Within 30 seconds, aim at the target chest and repeat the exact same gesture.")
numbered(doc, "The HUD confirms “Linked as Input/Fuel Input/Output/Replenishment for <station name>.”")
numbered(doc, "To remove, hold Shift with the same Alt+mouse gesture at both the station and linked chest.")
control_table(doc, [
    ("Alt+Left", "Input; native fuel control starts Fuel Input."),
    ("Alt+Right", "Output."), ("Alt+Middle", "Replenishment."),
    ("Shift+Alt+same mouse button", "Remove that role link."),
])
h2(doc, "Rings")
para(doc, "Green identifies Input and Fuel Input, yellow identifies Output, and turquoise identifies Replenishment. Focusing a station displays existing/empty roles and shows linked chest rings.")

page_break(doc)
title(doc, "Runic Production — station matrix", "production_matrix", "Detailed mod reference")
nav(doc)
kv_table(doc, [
    ("Smelter", "Input ore • Fuel Input coal • Output bars."),
    ("Cooking station / oven", "Input cookables • Fuel Input if the native station has fuel • Output cooked items • optional Replenishment."),
    ("Recipe station / mead ketill", "Input ingredients • Output products • Replenishment exemplar selects desired recipe."),
    ("Fermenter", "Input mead base • Output finished mead • optional Replenishment."),
    ("Fire / lamp", "Fuel Input from a linked wood/resin chest when the native fuel control is targeted."),
])
h2(doc, "Replenishment semantics")
bullet(doc, "Place a physical exemplar of the item to maintain in each replenishment chest."),
bullet(doc, "If the exemplar count falls below its reserve, the station may pull required ingredients from eligible nearby chests and make more."),
bullet(doc, "An empty replenishment chest does nothing; an Output link alone does not choose a recipe."),
bullet(doc, "A relay chest can be output for the ketill and input for the fermenter. This is a supported multi-role chain."),
h2(doc, "Authority and failure behavior")
para(doc, "Automation requires local ownership, native access, loaded/in-range objects, compatible station state, and a provable item transaction. If state becomes indeterminate, the station may be quarantined for that process session: Runic automation pauses and suppresses a risky fallback until the owning process reloads. This protects against uncertain duplication or loss, not just smelting/cooking loss.")
callout(doc, "Restart after session quarantine", "Reload the world for a client-owned/single-player station, or restart the dedicated server when it owns the automation. Then correct the underlying error before relying on the line.", PALE_RED, RED)

module_header(doc, "Runic Safety", "1.0.0", "Repeat-to-confirm guards around destructive or high-value actions.", "mod_safety")
h2(doc, "What it does", "safety_confirmations")
para(doc, "Runic Safety intercepts a small set of high-impact actions and asks for the identical action again within the confirmation window (default 4 seconds). It also denies operations that would move protected equipped, quest, or Runic Inventory-locked items through an unsafe path.")
h2(doc, "Protected actions")
bullet(doc, "Removing an occupied container."),
bullet(doc, "Removing a ship or cart."),
bullet(doc, "Overwriting a portal route/tag where supported."),
bullet(doc, "Rare sacrifice or similar high-value actions when the native target can be proven."),
h2(doc, "Other safeguards")
bullet(doc, "A native tombstone audit reports recovery state without inventing a second tombstone store."),
bullet(doc, "A migration-backup API is available to supported modules; Safety itself is not a general journal or rollback engine."),
bullet(doc, "If the target changes, the confirmation does not carry over. Repeat the same action on the same target."),
h2(doc, "Mini-base use")
para(doc, "Safety is most visible when dismantling a stocked chest, ship, cart, or configured portal. Ordinary building and transfers remain unobstructed.")
callout(doc, "Verify", "Attempt to dismantle a non-empty test chest. The first action should warn; repeating the same action within the window should allow the native result.", PALE_GREEN, GREEN)

module_header(doc, "Runic Sentinel", "1.0.0", "Signed mod-policy validation for administrators.", "mod_sentinel")
h2(doc, "What it does")
para(doc, "Runic Sentinel hashes the loaded plugin set and validates a signed RUNIC-SENTINEL/2 policy with a pinned RSA public key. It can compare the local state to an administrator-approved policy without transmitting world data or holding the private signing key.")
h2(doc, "Operating modes")
kv_table(doc, [
    ("Disabled", "Sentinel is inert."),
    ("Optional", "Default monitor-only posture; report mismatches without requiring disconnect."),
    ("Required", "After authentication, disconnect clients that do not satisfy the valid signed policy."),
])
h2(doc, "Required-mode checklist")
numbered(doc, "Create and sign the RUNIC-SENTINEL/2 policy outside the game with the private key kept offline."),
numbered(doc, "Distribute the identical policy and pinned public key to the server and every participating client."),
numbered(doc, "Validate in Optional mode first. Move to Required only after clean matches are proven."),
h2(doc, "Boundaries")
para(doc, "Sentinel is not an anti-cheat scanner, telemetry client, updater, or world-state validator. It verifies the declared plugin policy and signature. A bad or missing policy must not silently broaden authority.")
callout(doc, "Mini-base role", "Sentinel does not change the base. It verifies that the processes operating it are running the intended approved mod set.", PALE_BLUE, BLUE)

module_header(doc, "Runic Storage", "1.0.0", "Safe chest discovery, transfer, sorting, restock, consolidation, and search.", "mod_storage")
h2(doc, "What it does")
para(doc, "Runic Storage adds authorized hover summaries and bounded container actions around the local player. It uses native local ownership and access checks, does not claim remote chests, and does not create an RPC transfer service or durable transaction journal.")
h2(doc, "Default controls", "storage_controls")
control_table(doc, [
    ("Alt+Q", "Quick Stack carried items into eligible nearby matching stacks."),
    ("Alt+R", "Restock configured carried targets from eligible nearby chests."),
    ("Alt+F", "Open nearby-item search."),
    ("Alt+S", "Sort the opened container."),
    ("Alt+A", "Store all eligible carried items into the opened container."),
    ("Alt+C", "Consolidate carried stacks."),
])
h2(doc, "Protection integration")
para(doc, "When Runic Inventory is present, locked slots, role slots, equipped items, and other protected carried items are excluded from unsafe store/sort/consolidation paths. If protection cannot be proven, the action fails closed with no mutation.")
callout(doc, "Verify", "With a chest open, lock one carried slot and use Store All. The locked item remains; eligible items move. Close the chest, use Alt+Q, and confirm only matching nearby stacks receive items.", PALE_GREEN, GREEN)

page_break(doc)
title(doc, "Runic Storage — Alt+F search", "storage_search", "Detailed mod reference")
nav(doc)
h2(doc, "Search workflow")
numbered(doc, "Press Alt+F near loaded chests. The native Valheim-styled nearby-items window takes mouse and keyboard focus."),
numbered(doc, "Type in the filter or click an item row. Mouse clicks inside the window are consumed so the player does not punch or swing a weapon."),
numbered(doc, "Selecting an item marks every nearby chest containing it with a bright yellow ring/light for about 15 seconds."),
numbered(doc, "Close the window or choose another item. The highlight is discovery only; it does not transfer the item."),
h2(doc, "Appearance settings")
para(doc, "Menu font size and font color are configurable. Color is selected from a dropdown of actual color names rather than requiring a hexadecimal code. Use a high-contrast color appropriate to the Valheim panel skin.")
h2(doc, "Search boundaries")
bullet(doc, "Only eligible loaded chests inside the configured range are listed."),
bullet(doc, "Private/warded or otherwise unauthorized contents are not exposed."),
bullet(doc, "The menu uses native-like panels and controls, but it remains a Runic Storage search window—not a replacement inventory."),
callout(doc, "Troubleshoot", "If the cursor cannot click or typing does not enter the filter, verify the current package is loaded and inspect the client log for StorageSearch UI exceptions.", PALE_RED, RED)

module_header(doc, "Runic Velocity", "1.0.0", "Integrity-checked startup manifest caching and milestone diagnostics.", "mod_velocity")
h2(doc, "What it does")
para(doc, "Runic Velocity builds and validates a local manifest of plugin files so unchanged startup scans can reuse proven metadata. Background work inspects files without loading plugin code. It also records bounded startup milestones to help identify slow phases.")
h2(doc, "What the cache means")
bullet(doc, "A warm-cache result is a performance hint, not authorization to trust or execute a plugin."),
bullet(doc, "Changed size, timestamp, or integrity data invalidates the cached record and triggers fresh work."),
bullet(doc, "The mod does not modify another plugin, skip BepInEx dependency checks, or hide load failures."),
h2(doc, "Controls and configuration")
para(doc, "There is no gameplay UI or hotkey. Configuration is read at startup, so restart after changing Velocity settings. Useful evidence appears in the BepInEx log.")
h2(doc, "Mini-base role")
para(doc, "Velocity helps the modded profile reach the world efficiently. It has no effect on chests, stations, portals, crops, or building once play begins.")
callout(doc, "Verify", "Compare two consecutive clean launches. The second may report a valid warm manifest, while every plugin still loads and reports normally.", PALE_GREEN, GREEN)

module_header(doc, "Runic World Engine", "1.0.0", "Observe-only aggregate world and network diagnostics.", "mod_worldengine")
h2(doc, "What it does")
para(doc, "Runic World Engine samples aggregate ZDO population, peer activity, create/destroy events, send/receive work, and save/load durations. Sampling is bounded to at most once per second, and periodic summaries are off by default.")
h2(doc, "What it does not do")
bullet(doc, "No world mutation, registry replacement, global object scan, or unknown-data deletion."),
bullet(doc, "No player-facing dashboard is required. The primary output is diagnostic logging when enabled."),
bullet(doc, "Unknown world data remains owned by Valheim and the mod that created it."),
h2(doc, "Mini-base role")
para(doc, "Use World Engine when checking whether a busy production/storage base correlates with unusual aggregate object or network activity. It observes the operating environment; it does not tune or automate the base.")
h2(doc, "Configuration")
para(doc, "Enable summaries only when gathering evidence, keep the sampling interval bounded, and return verbose diagnostics to normal after the test. Compare trends rather than treating a single sample as a failure.")
callout(doc, "Verify", "Enable a short diagnostic summary, load the base, move through its active area, and confirm bounded aggregate lines appear without gameplay changes.", PALE_GREEN, GREEN)

# Quick glossary / configuration
page_break(doc)
title(doc, "Key concepts and where to find them", "key_concepts", "Linked reference")
for label, anchor, desc in (
    ("Eligible nearby chest", "mod_crafting", "Access-checked, loaded, in-range material source; Crafting explains the common model."),
    ("Agriculture ghost colors", "agriculture_ghosts", "Green ready, red resource shortage, amber invalid placement."),
    ("Agriculture fill order", "agriculture_fill", "Right-to-left grids and center-out shapes."),
    ("Protected inventory row", "inventory_roles", "Five gear roles and three quick slots in the native bottom row."),
    ("Alt+F search", "storage_search", "Keyboard/mouse-focused native-style item search and chest highlights."),
    ("Production link", "production_links", "Two-click station-to-chest relationship with a 30-second window."),
    ("Relay chest", "production_matrix", "Output for one station and input for another."),
    ("Replenishment exemplar", "production_matrix", "Physical item that selects what the station should maintain."),
    ("Portal network", "portal_modes", "Named permission-aware endpoints, distinct from Standard Pair."),
    ("Last-known pin", "exploration_tags", "Saved map information, not live entity tracking."),
    ("Local-axis rotation", "precision_controls", "Pitch/roll/yaw relative to the selected piece."),
    ("Session quarantine", "production_matrix", "Automation pause after indeterminate station state."),
):
    content_link_paragraph(doc, label, anchor, desc)
h2(doc, "Configuration files")
para(doc, "Most packages use BepInEx/config/chazman.<RunicModName>.cfg, derived from the plugin GUID. Runic Display Stands may use its package-specific GUID/config naming. Prefer Configuration Manager for discoverable runtime settings; use the files for deployment review and startup-only options.")
callout(doc, "Version scope", "This primer describes the current canonical packages: fourteen version 1.0.0 mods, Runic Display Stands 1.3.1, and Runic Precision Build Tool 2.0.1.", PALE_BLUE, BLUE)

# Package roster
page_break(doc)
title(doc, "Current package roster", "package_roster", "Reference")
roster = [
    ("Runic Agriculture", "1.0.0"), ("Runic Awareness", "1.0.0"), ("Runic Build Camera", "1.0.0"),
    ("Runic Crafting", "1.0.0"), ("Runic Display Stands", "1.3.1"), ("Runic Exploration", "1.0.0"),
    ("Runic Interaction", "1.0.0"), ("Runic Inventory", "1.0.0"), ("Runic Portals", "1.0.0"),
    ("Runic Precision Build Tool", "2.0.1"), ("Runic Production", "1.0.0"), ("Runic Safety", "1.0.0"),
    ("Runic Sentinel", "1.0.0"), ("Runic Storage", "1.0.0"), ("Runic Velocity", "1.0.0"),
    ("Runic World Engine", "1.0.0"),
]
table = doc.add_table(rows=9, cols=2)
table.style = "Table Grid"
fixed_table(table, [4680, 4680])
for col, heading_text in enumerate(("Package / version", "Package / version")):
    shade(table.cell(0, col), PALE_BLUE)
    set_font(table.cell(0, col).paragraphs[0].add_run(heading_text), 9.5, True, INK)
repeat_header(table.rows[0])
for idx, (package, version) in enumerate(roster):
    row = 1 + (idx % 8)
    col = idx // 8
    if row % 2:
        shade(table.cell(row, col), PALE_GREY)
    p = table.cell(row, col).paragraphs[0]
    set_font(p.add_run(package + "  "), 9.2, True, DARK_BLUE)
    set_font(p.add_run(version), 9.2, False, BODY)
h2(doc, "Installation boundary")
para(doc, "These are standalone packages with BepInEx dependencies. The retired Runic Core, Persistence, Permissions, Transactions, and Integrity packages are not part of this current primer or the mini-base setup.")
h2(doc, "Suggested test order")
numbered(doc, "Profile/startup: Sentinel, Velocity, World Engine."),
numbered(doc, "Build: Awareness, Build Camera, Precision Build Tool, Crafting, Safety."),
numbered(doc, "Inventory/storage: Inventory, Storage, Interaction, Display Stands."),
numbered(doc, "Gameplay systems: Agriculture, Production, Portals, Exploration."),
callout(doc, "Document navigation", "Use the Contents page to jump to any phase or mod. Every detailed mod page provides links back to Contents and the mini-base walkthrough.", PALE_GOLD, GOLD)

# Closing
page_break(doc)
title(doc, "Mini-base completion checklist", "completion_checklist", "Walkthrough closeout")
checks = [
    "All 16 current packages load without an exception.",
    "Build Camera toggles in hammer use and leaves the avatar vulnerable.",
    "Precision Build pitches an already-rotated beam around its local axis.",
    "The Inventory bottom row labels, locks, Quick slots, and repair behavior work.",
    "Alt+F accepts typing/clicks without attacking and highlights matching chests.",
    "The agriculture grid dimensions match the selected rows/columns; one left-click plants the ready cells.",
    "Resource shortage ghosts are red; invalid positions are amber.",
    "Crafting and building consume carried materials before exact nearby deficits.",
    "The ketill relay chest is both ketill Output and fermenter Input.",
    "Input/Fuel Input rings are green, Output yellow, and Replenishment turquoise.",
    "Fires/lamps use Fuel Input rather than an ambiguous ordinary Input label.",
    "Portal network permissions allow and deny the expected players.",
    "Exploration results are saved, explored, last-known pins only.",
    "Safety requires an intentional repeat for protected destructive actions.",
    "Background diagnostic mods remain non-invasive during normal play.",
]
for item in checks:
    p = doc.add_paragraph(style="List Bullet")
    set_font(p.add_run("☐  "), 11, True, BLUE, font="DejaVu Sans")
    p.add_run(item)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(14)
internal_link(p, "Return to Contents", "contents", DARK_BLUE, True)
p.add_run("   •   ")
internal_link(p, "Return to mini-base walkthrough", "walkthrough", DARK_BLUE, True)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
