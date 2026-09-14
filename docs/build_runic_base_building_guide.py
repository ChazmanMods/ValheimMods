from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\Valheim Mods\ChazmanModsRepo")
OUT = ROOT / "docs" / "Runic_Mod_Suite_Base_Building_Guide.docx"
SUITE_ICON = ROOT / "RunicModClientSuite" / "icon.png"
BUILD_IMAGE = ROOT / "RunicPrecisionBuildTool" / "media" / "runic_precision_angled_build.png"

BLACK = RGBColor(0, 0, 0)
BODY = RGBColor(39, 44, 51)
MUTED = RGBColor(94, 99, 106)
NAVY = "17324D"
PALE_BLUE = "EAF2F8"
PALE_GREY = "F5F6F7"
WHITE = RGBColor(255, 255, 255)
GOLD = RGBColor(165, 113, 32)


def set_run_font(run, size=None, bold=None, color=None, italic=None, name="Calibri"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=105, start=125, bottom=105, end=125):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        node = OxmlElement("w:tblHeader")
        node.set(qn("w:val"), "true")
        tr_pr.append(node)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_widths(table, widths_twips):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_twips)))
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_twips:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            width = widths_twips[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    shown = OxmlElement("w:t")
    shown.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, shown, end))


def add_picture(paragraph, path, width, alt_text):
    run = paragraph.add_run()
    inline = run.add_picture(str(path), width=Inches(width))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", alt_text)
    doc_pr.set("title", alt_text)
    return inline


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.34)
    section.footer_distance = Inches(0.36)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = BODY
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.14

    for style_name, size, before, after in (
        ("Title", 27, 0, 8),
        ("Subtitle", 13.5, 0, 11),
        ("Heading 1", 17, 10, 8),
        ("Heading 2", 12.7, 9, 4),
        ("Heading 3", 11.2, 7, 3),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.font.bold = style_name != "Subtitle"
        if style_name == "Subtitle":
            style.font.italic = False
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(10.4)
        style.font.color.rgb = BODY
        style.paragraph_format.left_indent = Inches(0.3)
        style.paragraph_format.first_line_indent = Inches(-0.17)
        style.paragraph_format.space_after = Pt(3.5)
        style.paragraph_format.line_spacing = 1.1

    if "Small Body" not in doc.styles:
        small = doc.styles.add_style("Small Body", WD_STYLE_TYPE.PARAGRAPH)
    else:
        small = doc.styles["Small Body"]
    small.font.name = "Calibri"
    small._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    small._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    small.font.size = Pt(9.1)
    small.font.color.rgb = BODY
    small.paragraph_format.space_after = Pt(2.5)
    small.paragraph_format.line_spacing = 1.04

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run("RUNIC MOD SUITE BASE BUILDING GUIDE"), 7.8, True, BLACK)

    first_hp = section.first_page_header.paragraphs[0]
    first_hp.text = ""

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    set_run_font(fp.add_run("Chazman Mods   Page "), 8, False, MUTED)
    add_page_number(fp)

    first_fp = section.first_page_footer.paragraphs[0]
    first_fp.text = ""

    props = doc.core_properties
    props.title = "Runic Mod Suite Base Building Guide"
    props.subject = "A concise Ravenhold base walkthrough and player manual for all sixteen Runic systems"
    props.author = "Chazman Mods"
    props.keywords = "Valheim, Runic, base building, walkthrough, mod manual"
    props.comments = "Prepared for the current Runic client and server suites"


def paragraph(doc, text="", style=None, align=None, keep=False):
    p = doc.add_paragraph(style=style)
    if text:
        p.add_run(text)
    if align is not None:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_together = True
    return p


def rich_paragraph(doc, parts, style=None, align=None, keep=False):
    p = doc.add_paragraph(style=style)
    for text, bold, color, italic in parts:
        run = p.add_run(text)
        set_run_font(run, bold=bold, color=color, italic=italic)
    if align is not None:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_together = True
    return p


def heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def bullet(doc, text):
    return paragraph(doc, text, style="List Bullet")


def page_break(doc):
    doc.add_page_break()


def add_role_table(doc):
    rows = [
        ("Solo player", "Runic Mod Client Suite"),
        ("Remote player", "Runic Mod Client Suite"),
        ("Dedicated server", "Runic Mod Server Suite"),
        ("Listen host who also plays", "Both suites"),
        ("Remote Sentinel administrator", "Client Suite plus full Runic Sentinel"),
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_widths(table, [2800, 6900])
    set_table_borders(table)
    for index, label in enumerate(("Play style", "Install")):
        cell = table.cell(0, index)
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(p.add_run(label), 9.5, True, WHITE)
    repeat_header(table.rows[0])
    for row_index, (profile, install) in enumerate(rows, start=1):
        cells = table.add_row().cells
        if row_index % 2 == 0:
            for cell in cells:
                set_cell_shading(cell, PALE_BLUE)
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
        p = cells[0].paragraphs[0]
        set_run_font(p.add_run(profile), 9.4, True, BODY)
        p = cells[1].paragraphs[0]
        set_run_font(p.add_run(install), 9.4, False, BODY)
        prevent_row_split(table.rows[row_index])
    set_table_widths(table, [2800, 6900])
    return table


def add_field_guide_table(doc, records):
    table = doc.add_table(rows=1, cols=3)
    set_table_widths(table, [2140, 3200, 4360])
    set_table_borders(table)
    for index, label in enumerate(("Runic system", "Best feature", "How to use it")):
        cell = table.cell(0, index)
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        set_run_font(p.add_run(label), 9.1, True, WHITE)
    repeat_header(table.rows[0])
    for row_index, (name, feature, use) in enumerate(records, start=1):
        cells = table.add_row().cells
        if row_index % 2 == 0:
            for cell in cells:
                set_cell_shading(cell, PALE_GREY)
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=90, bottom=90)
        p = cells[0].paragraphs[0]
        p.style = doc.styles["Small Body"]
        set_run_font(p.add_run(name), 9.2, True, BLACK)
        for col_index, value in ((1, feature), (2, use)):
            p = cells[col_index].paragraphs[0]
            p.style = doc.styles["Small Body"]
            set_run_font(p.add_run(value), 9.0, False, BODY)
        prevent_row_split(table.rows[row_index])
    set_table_widths(table, [2140, 3200, 4360])
    return table


doc = Document()
configure_document(doc)

# Cover and quick setup
p = paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
p.paragraph_format.space_before = Pt(5)
p.paragraph_format.space_after = Pt(7)
add_picture(p, SUITE_ICON, 1.72, "Runic Mod Client Suite emblem")

p = paragraph(doc, "Runic Mod Suite Base Building Guide", style="Title", align=WD_ALIGN_PARAGRAPH.CENTER)
p.paragraph_format.space_before = Pt(0)
title_ppr = p._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
p = paragraph(doc, "Ravenhold Walkthrough and Player Manual", style="Subtitle", align=WD_ALIGN_PARAGRAPH.CENTER)
set_run_font(p.runs[0], 13.5, False, BLACK)

p = paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
p.paragraph_format.space_after = Pt(13)
set_run_font(p.add_run(date.today().strftime("%B %Y edition")), 9.4, False, MUTED)

paragraph(
    doc,
    "This guide turns all 16 Runic systems into one practical build: Ravenhold, a compact base with a great hall, connected workshop, farm, automated kitchen and forge, armory, storage hall, and portal tower. Follow the walkthrough once, then use the two-page field guide when you need a key or feature reminder.",
    keep=True,
)
paragraph(
    doc,
    "Normal building costs, wards, and access rules still apply. The suite adds clearer information and faster controls while keeping the base inside Valheim's normal rules.",
    keep=True,
)

heading(doc, "Choose the Right Suite", 1)
add_role_table(doc)

heading(doc, "The Three Keys to Start", 2)
bullet(doc, "B toggles Runic Build Camera while a build tool is equipped.")
bullet(doc, "P toggles Runic Precision Build Tool during that hammer build session.")
bullet(doc, "Left Alt is the suite's usual modifier for storage, farming, production links, quick slots, and protected actions.")

page_break(doc)

# Walkthrough page one
heading(doc, "Build Ravenhold", 1)
paragraph(
    doc,
    "Ravenhold is small enough for one player and complete enough for a group. Keep the workshop, storage hall, kitchen, and forge near the center; put the farm outside one wall and raise the portal tower above the gate.",
)

heading(doc, "1 Choose the Site", 2)
paragraph(
    doc,
    "Open Valheim's large map and use Runic Exploration to search places you have already discovered. Pick a coast or river bend with room for a field and dock. Add useful pin names such as [bed] Ravenhold, [portal] Raven Gate, and [boat] Longship. The browser can find those tags later and center the map on the result. It never reveals fog or live-tracks an object.",
)
paragraph(
    doc,
    "As you walk the site, Runic Awareness quietly shows current food, effects, comfort, nearby build details, station status, crop growth, and tameable context. There is no key to remember; the panel hides when a menu or map needs the screen.",
)

heading(doc, "2 Raise the Great Hall", 2)
paragraph(
    doc,
    "Equip the hammer and press B. Fly the camera with normal movement and look controls; Jump rises, Crouch descends, and Run moves faster. Your Viking stays where you left them and remains vulnerable, while building costs, wards, stations, support, and distance limits still apply.",
)
paragraph(
    doc,
    "Press P for exact placement. Use the wheel for yaw, Alt plus wheel for pitch, Shift plus wheel for roll, and add V for fine steps. Hold G to see local axes. F11 can undo the last eligible placement, and F12 repairs a small local group of pieces. Build one dramatic angled doorway or roof brace, then leave the rest of the hall simple.",
)
p = paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
p.paragraph_format.space_before = Pt(5)
p.paragraph_format.space_after = Pt(1)
add_picture(p, BUILD_IMAGE, 3.15, "Angled Valheim timber construction using Runic Precision Build Tool")
p = paragraph(doc, "A single precise feature gives the hall character without slowing the whole build", align=WD_ALIGN_PARAGRAPH.CENTER)
set_run_font(p.runs[0], 8.6, False, MUTED, True)

page_break(doc)

# Walkthrough page two
heading(doc, "Build the Workshop Farm and Armory", 1)

heading(doc, "3 Connect the Workshop", 2)
paragraph(
    doc,
    "Place shared supply chests around the workbench, forge, and main build area. Keep them within about 20 metres. Runic Crafting uses carried materials first, then takes the remaining materials from nearby accessible chests. Requirement rows show the combined total, and one press of the normal repair button repairs every available worn item.",
)
paragraph(
    doc,
    "Seed each category chest with one example item. After a trip, Runic Storage can now recognize where matching stacks belong. Hover a closed chest for a quick preview, use Alt+Q to quick stack into nearby seeded chests, and use Alt+F when you want the right chest to glow.",
)

heading(doc, "4 Build the Armory", 2)
paragraph(
    doc,
    "Runic Inventory turns the existing bottom row into Helmet, Chest, Legs, Cape, Utility, Q1, Q2, and Q3. It adds no capacity. Use Alt+1, Alt+2, and Alt+3 outside the inventory for quick-slot items, Alt+I to sort ordinary rows, and Alt plus right-click to lock an important slot.",
)
paragraph(
    doc,
    "Place item stands and armor stands by the door. Interact to open one like a container, arrange the kit, then choose Use/equip to equip an item or swap the supported armor loadout. Runic Interaction makes the room feel faster: hold Use to repeat supported feeding actions and Alt-click to transfer a full stack.",
)

heading(doc, "5 Plant the Courtyard Farm", 2)
paragraph(
    doc,
    "Equip the cultivator and select a crop. Press Alt+O to cycle Row, Grid, Circle, Star, triangle, half-circle, and trapezoid patterns. Rotate with the wheel; Alt plus wheel changes rows, Shift plus wheel changes columns, and Alt+Shift plus wheel changes spacing. Green previews are ready, amber previews fail a placement rule, and red previews need more seed.",
)
paragraph(
    doc,
    "Plant normally when the pattern looks right. When the crop matures, Alt+E harvests a small area. Select the same crop and use Alt+T to confirm replanting at the saved positions. Nearby seed chests can supply the job when access and ownership allow it.",
)

page_break(doc)

# Walkthrough page three
heading(doc, "Add Production Portals and Storage", 1)

heading(doc, "6 Automate Fire and Feast", 2)
paragraph(
    doc,
    "Runic Production links a station to nearby chests with the same gesture at both ends within 30 seconds. Alt+Left links Input. Aim at the station's fuel opening and use Alt+Left for Fuel Input. Alt+Right links Output. Alt+Middle links Replenishment. Add Shift at both ends to unlink.",
)
paragraph(
    doc,
    "Start with a clean smelter line: ore chest to Input, coal chest to Fuel Input, and bar chest to Output. Then build the showpiece mead line. Link an ingredient chest to the recipe station, place one desired mead base in a Replenishment chest, use that chest as fermenter Input, and place one finished mead in a fermenter Replenishment chest. Green rings mark inputs and fuel, yellow marks output, and turquoise marks replenishment.",
)
paragraph(
    doc,
    "Keep the example item inside every Replenishment chest. Production pauses whenever a linked station or chest is unavailable and resumes when it is ready again.",
)

heading(doc, "7 Open the Portal Tower", 2)
paragraph(
    doc,
    "Aim at a standard portal and press Left Shift+E to open the Runic editor. Create the home gate with network|private|Ravenhold|Great Hall|both. Use a unique endpoint name for each destination; network spelling is case-sensitive. Walk into a depart or both portal, choose an authorized destination on the map, and travel. Press P on the normal large map to toggle the portal directory.",
)
paragraph(
    doc,
    "Public networks still respect wards, private networks are owner-only, and group networks follow the active group. Runic Safety asks for confirmation before an important portal is changed.",
)

heading(doc, "8 Finish the Raid Loop", 2)
paragraph(
    doc,
    "Return through the gate and make unloading one smooth circuit. Hover chests to preview them, press Alt+Q to send matching items home, open an overflow chest and press Alt+A for eligible leftovers, then Alt+C to consolidate partial carried stacks. Alt+R restores your configured travel stock, and Alt+F finds the chest holding the missing item.",
)
paragraph(
    doc,
    "Runic Safety asks for confirmation before risky actions. Runic Velocity helps diagnose startup problems, and Runic World Engine helps keep world saves smooth. All three work automatically during normal play.",
)

page_break(doc)

# Field guide page one
heading(doc, "Runic Field Guide", 1)
paragraph(doc, "Use this section when you remember the feature but not the control. Defaults can be changed in the BepInEx configuration files.")

records_one = [
    (
        "Runic Agriculture",
        "Pattern planting, nearby seed use, area harvest, confirmed replant, and crop or beehive status.",
        "With the cultivator active: Alt+O cycles shapes; wheel rotates; Alt or Shift with wheel changes the layout; Alt+E harvests; Alt+T replants.",
    ),
    (
        "Runic Awareness",
        "Automatic food, effects, comfort, building, station, crop, and tameable context.",
        "No hotkey. Look at the world and play normally. The display hides during inventory, map, chat, console, and other menus.",
    ),
    (
        "Runic Build Camera",
        "Detached camera placement from better angles while the player remains in place and vulnerable.",
        "Equip a build tool and press B. Use normal move and look controls; Jump and Crouch move vertically; Run accelerates.",
    ),
    (
        "Runic Crafting",
        "Crafting and building from nearby accessible chests, combined requirement counts, and Repair All.",
        "Use the normal craft, build, and repair controls. Carried materials are used first, then the remaining materials come from nearby chests.",
    ),
    (
        "Runic Display Stands",
        "Container-style item and armor stands for displays and quick loadout swaps.",
        "Interact to open the stand. Drag items normally. Choose Use/equip for a single item or supported armor loadout; Alt+Interact takes one item.",
    ),
    (
        "Runic Exploration",
        "Searchable, last-known pins plus direction, elevation, distance, and a sailing readout.",
        "Open the large map, search saved explored pins, filter, and click a result to center it. Helpful tags include [portal], [bed], [boat], and [tombstone].",
    ),
    (
        "Runic Interaction",
        "Repeated supported Use actions, full-stack transfer, pickup filtering, and small menu or equipment polish.",
        "Hold Use to repeat supported feeding. Alt-click transfers a full stack. Alt+Use bypasses an enabled pickup filter. Door auto-close is optional and off by default.",
    ),
    (
        "Runic Inventory",
        "Equipment roles and three quick slots in Valheim's existing bottom row, plus slot locks and sorting.",
        "Alt+1 through Alt+3 use Q1 through Q3. Alt+I sorts ordinary rows. Alt+right-click locks the pointed slot. No extra capacity is added.",
    ),
]
add_field_guide_table(doc, records_one)

page_break(doc)

# Field guide page two
heading(doc, "Runic Field Guide Continued", 1)
records_two = [
    (
        "Runic Portals",
        "Named public, private, or group networks with a destination map and authorized portal directory.",
        "Left Shift+E edits a standard portal. Example: network|private|Ravenhold|Great Hall|both. Walk in to choose a destination; P toggles the map directory.",
    ),
    (
        "Runic Precision Build Tool",
        "Six-axis placement, exact matching, local guides, last-placement undo, and nearby repair.",
        "Press P while building. Wheel, Alt+wheel, and Shift+wheel rotate; V makes fine steps; G shows axes; F11 undoes; F12 repairs nearby eligible pieces.",
    ),
    (
        "Runic Production",
        "Input, fuel, output, and replenishment chest links for supported stations.",
        "Repeat the same station-to-chest gesture within 30 seconds: Alt+Left Input or Fuel, Alt+Right Output, Alt+Middle Replenishment. Add Shift to unlink.",
    ),
    (
        "Runic Safety",
        "Guards important items and asks for confirmation before risky actions.",
        "No hotkey. Follow the on-screen prompt and repeat the same action when confirmation is required. Move or unlock a protected item before using it.",
    ),
    (
        "Runic Sentinel",
        "Keeps multiplayer connections aligned with the server's mod rules.",
        "Sentinel Client works automatically for players. Approved administrators use full Sentinel and press F3 to open its panel.",
    ),
    (
        "Runic Storage",
        "Chest previews, Quick Stack, Restock, Search, Sort, Store All, and carried-stack consolidation.",
        "Alt+Q quick stacks; Alt+R restocks; Alt+F searches; Alt+S sorts an open chest; Alt+A stores eligible items; Alt+C consolidates carried stacks.",
    ),
    (
        "Runic Velocity",
        "Tracks mod startup so problems are easier to diagnose.",
        "Works automatically. Check the game log only when troubleshooting startup.",
    ),
    (
        "Runic World Engine",
        "Helps the server handle world saves smoothly and records useful performance information.",
        "Works automatically in the background.",
    ),
]
add_field_guide_table(doc, records_two)

heading(doc, "Controller Note", 2)
paragraph(
    doc,
    "Agriculture, Build Camera, Interaction, Inventory, and Storage include controller controls. Runic Precision Build Tool is keyboard and mouse only. Check each mod's configuration if your controller buttons are remapped.",
)

# Shared-world setup
heading(doc, "Run a Shared World", 1)
paragraph(
    doc,
    "Dedicated servers use Runic Mod Server Suite, while players use Runic Mod Client Suite. Sentinel Client works automatically for players. Server owners and approved administrators use full Runic Sentinel and press F3 to open its panel. Detailed policy setup is covered in the Runic Sentinel README.",
)

heading(doc, "Complete the Ravenhold Tour", 2)
paragraph(
    doc,
    "Take one final lap: launch Build Camera from the gate, inspect the great hall, swap a loadout, harvest the courtyard, watch the forge and mead line move items, open the portal map, then return from a short raid and unload with Storage. When every step works, all 16 Runic systems are active in one base.",
)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
